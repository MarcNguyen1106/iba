"""Build BUSM1042_A3_Q2_LeadTime.xlsx (Assignment 3, Q2: lead time and cancellation).

Usage:  python3 -I build_workbook.py <source_bookings.xlsx> <output.xlsx> [--values-fallback]

The source file is only read, never written. Every summary number in the workbook is an
Excel formula; the only Python-computed column is Clean_Data!dup_flag (see README sheet).
"""
import re
import sys
import time
from datetime import datetime

import pandas as pd
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, Series
from openpyxl.chart.data_source import AxDataSource, StrRef
from openpyxl.chart.label import DataLabelList
from openpyxl.chart.series import SeriesLabel
from openpyxl.chart.shapes import GraphicalProperties
from openpyxl.chart.text import RichText, Text
from openpyxl.chart.title import Title
from openpyxl.comments import Comment
from openpyxl.drawing.line import LineProperties
from openpyxl.drawing.text import (CharacterProperties, Font as DFont, Paragraph,
                                   ParagraphProperties, RegularTextRun, RichTextProperties)
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties
from openpyxl.worksheet.table import Table, TableStyleInfo

SRC, OUT = sys.argv[1], sys.argv[2]
VALUES_FALLBACK = "--values-fallback" in sys.argv
BUILD_DATE = "23 October 2026"
N = 119390
LAST = N + 1  # last data row on Raw_Data / Clean_Data

# ---------------------------------------------------------------- style constants
FONT = "Calibri"
NAVY, GRID, BAND, TEAL, LGREY = "1F3A5F", "BFBFBF", "F2F2F2", "CCECE8", "D9D9D9"
F_BODY = Font(name=FONT, size=11)
F_BOLD = Font(name=FONT, size=11, bold=True)
F_TITLE = Font(name=FONT, size=14, bold=True, color=NAVY)
F_TTITLE = Font(name=FONT, size=12, bold=True, color=NAVY)
F_HDR = Font(name=FONT, size=11, bold=True, color="FFFFFF")
F_NOTE = Font(name=FONT, size=10, italic=True, color="595959")
F_SMALL = Font(name=FONT, size=10, color="404040")
FILL_HDR = PatternFill("solid", fgColor=NAVY)
FILL_BAND = PatternFill("solid", fgColor=BAND)
FILL_GREY = PatternFill("solid", fgColor=LGREY)
FILL_TEAL = PatternFill("solid", fgColor=TEAL)
THIN = Side(style="thin", color=GRID)
B_ALL = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
B_TOTAL = Border(left=THIN, right=THIN, top=Side(style="medium", color="404040"), bottom=THIN)
A_HDR = Alignment(horizontal="center", vertical="center", wrap_text=True)
A_WRAP = Alignment(vertical="top", wrap_text=True)
A_LEFT = Alignment(vertical="center")
A_NUM = Alignment(horizontal="right", vertical="center")
A_CTR = Alignment(horizontal="center", vertical="center")
FMT_N, FMT_P, FMT_RR = "#,##0", "0.0%", "0.00"
OK, BAD = "✔ OK", "✘ CHECK"
GREEN_FILL, GREEN_FONT = PatternFill("solid", fgColor="C6EFCE"), Font(name=FONT, color="006100", bold=True)
RED_FILL, RED_FONT = PatternFill("solid", fgColor="FFC7CE"), Font(name=FONT, color="9C0006", bold=True)

ENDASH = "–"
BANDS = [("0–7", 0, 7), ("8–30", 8, 30), ("31–90", 31, 90), ("91–180", 91, 180),
         ("181–365", 181, 365), ("Over 365", 366, 100000)]


def lit(text):
    """Excel string literal; long text is split because a literal may hold at most 255 chars."""
    chunks = [text[i:i + 200] for i in range(0, len(text), 200)] or [""]
    return "&".join('"' + c.replace('"', '""') + '"' for c in chunks)


def put(ws, ref, value, font=F_BODY, fmt=None, align=None, fill=None, border=None):
    c = ws[ref]
    c.value = value
    c.font = font
    if fmt:
        c.number_format = fmt
    if align:
        c.alignment = align
    if fill:
        c.fill = fill
    if border:
        c.border = border
    return c


def table_block(ws, top, left, headers, rows, fmts, total_rows=(), widths=None, hdr_height=None,
                wrap_cols=()):
    """Write a styled table: navy header row then body rows with banding. Returns last row."""
    for j, h in enumerate(headers):
        c = ws.cell(row=top, column=left + j, value=h)
        c.font, c.fill, c.alignment, c.border = F_HDR, FILL_HDR, A_HDR, B_ALL
    if hdr_height:
        ws.row_dimensions[top].height = hdr_height
    for i, row in enumerate(rows):
        r = top + 1 + i
        is_total = i in total_rows
        for j, v in enumerate(row):
            c = ws.cell(row=r, column=left + j, value=v)
            fmt = fmts[j] if j < len(fmts) else None
            c.font = F_BOLD if is_total else F_BODY
            if fmt and not (fmt == "@" and isinstance(v, str) and v.startswith("=")):
                c.number_format = fmt  # never text-format a formula cell (it would stop recalculating if edited)
            c.alignment = A_WRAP if j in wrap_cols else (A_LEFT if j == 0 or fmt == "@" else A_NUM)
            c.border = B_TOTAL if is_total else B_ALL
            if i % 2 == 1 and not is_total:
                c.fill = FILL_BAND
    if widths:
        for j, w in enumerate(widths):
            if w:
                ws.column_dimensions[get_column_letter(left + j)].width = w
    return top + len(rows)


def shown_len(v):
    """Approximate displayed length of a cell: for a formula, its text literals plus ~6 characters per reference."""
    v = str(v)
    if not v.startswith("="):
        return len(v)
    literals = re.findall(r'"((?:[^"]|"")*)"', v)
    return sum(len(x) for x in literals) + 6 * max(v.count("&") // 2, 0)


def title(ws, ref, text):
    put(ws, ref, text, font=F_TITLE)
    ws.row_dimensions[ws[ref].row].height = 22


def ttitle(ws, ref, text):
    put(ws, ref, text, font=F_TTITLE)
    ws.row_dimensions[ws[ref].row].height = 18


def note(ws, ref, text):
    put(ws, ref, text, font=F_NOTE)


def presentation(ws, tab, landscape=True, area=None):
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = tab
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.print_options.horizontalCentered = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    if area:
        ws.print_area = area


def match_rules(ws, rng):
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{OK}"'], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'{first}="{BAD}"'], fill=RED_FILL, font=RED_FONT))


def bool_rules(ws, rng):
    first = rng.split(":")[0].replace("$", "")
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{first}=TRUE"], fill=GREEN_FILL, font=GREEN_FONT))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f"{first}<>TRUE"], fill=RED_FILL, font=RED_FONT))


t0 = time.time()
# ---------------------------------------------------------------- 1. load source (read-only)
print("Loading source with pandas (for dup_flag) ...", flush=True)
df = pd.read_excel(SRC, sheet_name="hotels", engine="openpyxl")
assert df.shape == (N, 29), df.shape
cols = list(df.columns)
dup = df.duplicated(subset=cols[1:], keep="first").astype(int).tolist()
assert sum(dup) == 32239, sum(dup)

print("Loading source cell values with openpyxl ...", flush=True)
src_wb = load_workbook(SRC, read_only=True)
src_rows = list(src_wb["hotels"].iter_rows(values_only=True))
src_wb.close()
header = list(src_rows[0])
assert header == cols
data = src_rows[1:]
assert len(data) == N
# The source stores missing values as the text "NULL" (country, agent, company) and "NA" (children);
# pandas reads these as NaN. They are written as genuinely empty cells (never "NULL"/"NA"/"nan"/0).
NULL_TEXT = {"children": "NA", "country": "NULL", "agent": "NULL", "company": "NULL"}
null_idx = {cols.index(k): v for k, v in NULL_TEXT.items()}
clean_rows = []
for r in data:
    r = list(r)
    for i, token in null_idx.items():
        if r[i] == token:
            r[i] = None
    clean_rows.append(r)
del data, src_rows
assert [r[0] for r in clean_rows] == df["booking_order_id"].tolist()
del df

COL = {name: get_column_letter(i + 1) for i, name in enumerate(cols)}
DATE_COL = cols.index("arrival_date")

wb = Workbook()
wb.calculation.fullCalcOnLoad = True
wb.calculation.calcMode = "auto"
ws_readme = wb.active
ws_readme.title = "README"

# ---------------------------------------------------------------- 2. Raw_Data
print("Writing Raw_Data ...", flush=True)
ws_raw = wb.create_sheet("Raw_Data")
ws_raw.sheet_properties.tabColor = "2E75B6"


def write_data_sheet(ws, rows):
    ws.append(cols)
    for r in rows:
        ws.append(r)
    for row in ws.iter_rows(min_row=2, max_row=LAST, min_col=DATE_COL + 1, max_col=DATE_COL + 1):
        c = row[0]
        if isinstance(c.value, datetime):
            c.number_format = "dd/mm/yyyy"
    for j in range(1, 30):
        h = ws.cell(row=1, column=j)
        h.font, h.fill, h.alignment = F_BOLD, FILL_GREY, A_HDR
        ws.column_dimensions[get_column_letter(j)].width = max(11, min(30, len(str(h.value)) + 3))
    ws.column_dimensions["A"].width = 17
    ws.column_dimensions[COL["arrival_date"]].width = 17
    ws.freeze_panes = "A2"


write_data_sheet(ws_raw, clean_rows)
ws_raw.auto_filter.ref = f"A1:AC{LAST}"
ws_raw["A1"].comment = Comment("Original data, do not edit", "BUSM1042 group")
ws_raw.protection.sheet = True          # no password
ws_raw.protection.autoFilter = False    # allow filtering on the protected sheet
ws_raw.protection.sort = False

# ---------------------------------------------------------------- 3. Clean_Data
print("Writing Clean_Data ...", flush=True)
ws_cl = wb.create_sheet("Clean_Data")
ws_cl.sheet_properties.tabColor = "2E75B6"
C = COL  # column letters of the original fields
derived = [
    ("cancel_flag", lambda r: f'=IF({C["is_canceled"]}{r}="Yes",1,0)'),
    ("repeat_flag", lambda r: f'=IF({C["is_repeated_guest"]}{r}="Yes",1,0)'),
    ("lead_band", lambda r: (f'=IF({C["lead_time"]}{r}<=7,"0–7",IF({C["lead_time"]}{r}<=30,"8–30",'
                             f'IF({C["lead_time"]}{r}<=90,"31–90",IF({C["lead_time"]}{r}<=180,"91–180",'
                             f'IF({C["lead_time"]}{r}<=365,"181–365","Over 365")))))')),
    ("lead_band_order", lambda r: (f'=IF({C["lead_time"]}{r}<=7,1,IF({C["lead_time"]}{r}<=30,2,'
                                   f'IF({C["lead_time"]}{r}<=90,3,IF({C["lead_time"]}{r}<=180,4,'
                                   f'IF({C["lead_time"]}{r}<=365,5,6)))))')),
    ("lead_over_60", lambda r: f'=IF({C["lead_time"]}{r}>60,1,0)'),
    ("country_clean", lambda r: f'=IF({C["country"]}{r}="","Unknown",{C["country"]}{r})'),
    ("agent_clean", lambda r: f'=IF({C["agent"]}{r}="","None",{C["agent"]}{r})'),
    ("company_clean", lambda r: f'=IF({C["company"]}{r}="","None",{C["company"]}{r})'),
    ("arrival_date_valid", lambda r: f'=IF(ISNUMBER({C["arrival_date"]}{r}),1,0)'),
    ("guests_total", lambda r: f'={C["adults"]}{r}+N({C["children"]}{r})+{C["babies"]}{r}'),
]
FORMULA_KEEP = {"cancel_flag", "lead_band", "lead_band_order", "arrival_date_valid"}


def py_value(name, row):
    """Python value of a derived column (used only in the --values-fallback mode)."""
    v = dict(zip(cols, row))
    lt = v["lead_time"]
    order = 1 if lt <= 7 else 2 if lt <= 30 else 3 if lt <= 90 else 4 if lt <= 180 else 5 if lt <= 365 else 6
    return {
        "repeat_flag": 1 if v["is_repeated_guest"] == "Yes" else 0,
        "lead_over_60": 1 if lt > 60 else 0,
        "country_clean": "Unknown" if v["country"] is None else v["country"],
        "agent_clean": "None" if v["agent"] is None else v["agent"],
        "company_clean": "None" if v["company"] is None else v["company"],
        "guests_total": v["adults"] + (v["children"] or 0) + v["babies"],
        "lead_band_order": order,
    }[name]


dcols = [d[0] for d in derived] + ["dup_flag"]
DC = {name: get_column_letter(30 + i) for i, name in enumerate(dcols)}
ws_cl.append(cols + dcols)
for i, r in enumerate(clean_rows):
    rr = i + 2
    extra = []
    for name, f in derived:
        if VALUES_FALLBACK and name not in FORMULA_KEEP:
            extra.append(py_value(name, r))
        else:
            extra.append(f(rr))
    ws_cl.append(r + extra + [dup[i]])
for row in ws_cl.iter_rows(min_row=2, max_row=LAST, min_col=DATE_COL + 1, max_col=DATE_COL + 1):
    if isinstance(row[0].value, datetime):
        row[0].number_format = "dd/mm/yyyy"
for j in range(1, 30):
    h = ws_cl.cell(row=1, column=j)
    h.font, h.fill, h.alignment = F_BOLD, FILL_GREY, A_HDR
    ws_cl.column_dimensions[get_column_letter(j)].width = max(11, min(30, len(str(h.value)) + 3))
ws_cl.column_dimensions["A"].width = 17
ws_cl.column_dimensions[C["arrival_date"]].width = 17
for name, f in derived:
    h = ws_cl[f"{DC[name]}1"]
    h.font, h.fill, h.alignment = F_BOLD, FILL_TEAL, A_HDR
    ws_cl.column_dimensions[DC[name]].width = max(12, len(name) + 4)
    how = "formula" if not (VALUES_FALLBACK and name not in FORMULA_KEEP) else "value (formula shown for reference)"
    h.comment = Comment(f"Derived column, {how}. Row 2: {f(2)}", "BUSM1042 group", width=320, height=110)
h = ws_cl[f"{DC['dup_flag']}1"]
h.font, h.fill, h.alignment = F_BOLD, FILL_TEAL, A_HDR
ws_cl.column_dimensions[DC["dup_flag"]].width = 11
h.comment = Comment("VALUE computed in Python: 1 if this row is identical to an EARLIER row on all 28 columns "
                    "except booking_order_id (pandas duplicated(keep='first'), blanks treated as equal), else 0. "
                    "A 28-column duplicate test is impractical as a formula over 119,390 rows.",
                    "BUSM1042 group", width=320, height=140)
LASTCOL = DC["dup_flag"]
tbl = Table(displayName="tblClean", ref=f"A1:{LASTCOL}{LAST}")
tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True, showColumnStripes=False,
                                    showFirstColumn=False, showLastColumn=False)
ws_cl.add_table(tbl)
ws_cl.freeze_panes = "A2"
del clean_rows

# ---------------------------------------------------------------- named ranges
def rng(col, sheet="Clean_Data"):
    return f"{sheet}!${col}$2:${col}${LAST}"


NAMES = {
    "rng_id": rng("A"), "rng_lead": rng(C["lead_time"]), "rng_cancel": rng(DC["cancel_flag"]),
    "rng_deposit": rng(C["deposit_type"]), "rng_segment": rng(C["market_segment"]),
    "rng_channel": rng(C["distribution_channel"]), "rng_band": rng(DC["lead_band"]),
    "rng_dup": rng(DC["dup_flag"]), "rng_over60": rng(DC["lead_over_60"]),
    "rng_iscanceled": rng(C["is_canceled"]), "rng_valid": rng(DC["arrival_date_valid"]),
    "rng_guests": rng(DC["guests_total"]),
}
for k, v in NAMES.items():
    wb.defined_names[k] = DefinedName(k, attr_text=v)


def raw(field):
    return rng(C[field], "Raw_Data")


# ---------------------------------------------------------------- 4. Q2_Summary (built before Table 1 links to it)
print("Writing analysis sheets ...", flush=True)
ws_t1 = wb.create_sheet("Table1_DataQuality")
ws_q = wb.create_sheet("Q2_Summary")
ws_f = wb.create_sheet("Figure2")
ws_ch = wb.create_sheet("Checks")
ws_rf = wb.create_sheet("Report_Figures")

q = ws_q
title(q, "A1", "Q2: How does booking lead time affect the likelihood of cancellation?")
put(q, "A2", "Every number on this sheet is a live COUNTIFS or arithmetic formula on Clean_Data (all 119,390 "
             "bookings). A cancellation is is_canceled = Yes, which includes no-shows.", font=F_SMALL)

# Table A
TA_T, TA_H = 4, 5
TA0 = TA_H + 1                     # first band row
TA_TOT = TA0 + len(BANDS)          # total row
ttitle(q, f"A{TA_T}", "Table A: Cancellation by lead-time band")
hdrA = ["Lead time band (days)", "Lower bound (days)", "Upper bound (days)", "Bookings", "Cancelled",
        "Cancellation rate (all)", "No-deposit bookings", "No-deposit cancelled", "No-deposit rate",
        "No-deposit Online TA bookings", "No-deposit Online TA cancelled", "No-deposit Online TA rate",
        "Non Refund bookings", "Non Refund share of band", "Share of all bookings", "Share of all cancellations"]
rowsA = []
for i, (lab, lo, hi) in enumerate(BANDS):
    r = TA0 + i
    band = f'rng_lead,">="&$B{r},rng_lead,"<="&$C{r}'
    rowsA.append([
        lab, lo, hi,
        f"=COUNTIFS({band})", f"=COUNTIFS({band},rng_cancel,1)", f"=E{r}/D{r}",
        f'=COUNTIFS({band},rng_deposit,"No Deposit")', f'=COUNTIFS({band},rng_deposit,"No Deposit",rng_cancel,1)',
        f"=H{r}/G{r}",
        f'=COUNTIFS({band},rng_deposit,"No Deposit",rng_segment,"Online TA")',
        f'=COUNTIFS({band},rng_deposit,"No Deposit",rng_segment,"Online TA",rng_cancel,1)', f"=K{r}/J{r}",
        f'=COUNTIFS({band},rng_deposit,"Non Refund")', f"=M{r}/D{r}",
        f"=D{r}/$D${TA_TOT}", f"=E{r}/$E${TA_TOT}",
    ])
f0, f1 = TA0, TA_TOT - 1
rowsA.append(["Total", None, None, f"=SUM(D{f0}:D{f1})", f"=SUM(E{f0}:E{f1})", f"=E{TA_TOT}/D{TA_TOT}",
              f"=SUM(G{f0}:G{f1})", f"=SUM(H{f0}:H{f1})", f"=H{TA_TOT}/G{TA_TOT}",
              f"=SUM(J{f0}:J{f1})", f"=SUM(K{f0}:K{f1})", "—",
              f"=SUM(M{f0}:M{f1})", f"=M{TA_TOT}/D{TA_TOT}", f"=SUM(O{f0}:O{f1})", f"=SUM(P{f0}:P{f1})"])
fmtA = ["@", "0", "0", FMT_N, FMT_N, FMT_P, FMT_N, FMT_N, FMT_P, FMT_N, FMT_N, FMT_P, FMT_N, FMT_P, FMT_P, FMT_P]
table_block(q, TA_H, 1, hdrA, rowsA, fmtA, total_rows=(len(BANDS),),
            widths=[24, 17, 12.5] + [12.5] * 13, hdr_height=48)
q[f"B{TA_TOT}"].value = q[f"C{TA_TOT}"].value = None
q[f"L{TA_TOT}"].alignment = A_CTR
q[f"C{TA_TOT - 1}"].number_format = "#,##0"
note(q, f"A{TA_TOT + 1}", "Bands are counted directly from lead_time (inclusive bounds; 'Over 365' = 366 days or more). "
                          "Online TA = market_segment 'Online TA'. No-deposit = deposit_type 'No Deposit'. "
                          "Source: Clean_Data (n = 119,390).")

# Table B
TB_T = TA_TOT + 3
TB_H = TB_T + 1
ttitle(q, f"A{TB_T}", "Table B: 60-day comparison (D-EDGE benchmark)")
hdrB = ["Group", "Lead time", "Bookings", "Cancelled", "Cancellation rate", "Relative risk (>60 ÷ ≤60)",
        "% more likely to cancel"]
b1 = TB_H + 1
rowsB = [
    ["No-deposit bookings", "More than 60 days", '=COUNTIFS(rng_deposit,"No Deposit",rng_lead,">60")',
     '=COUNTIFS(rng_deposit,"No Deposit",rng_lead,">60",rng_cancel,1)', f"=D{b1}/C{b1}",
     f"=E{b1}/E{b1 + 1}", f"=F{b1}-1"],
    ["No-deposit bookings", "60 days or less", '=COUNTIFS(rng_deposit,"No Deposit",rng_lead,"<=60")',
     '=COUNTIFS(rng_deposit,"No Deposit",rng_lead,"<=60",rng_cancel,1)', f"=D{b1 + 1}/C{b1 + 1}",
     "reference", "reference"],
    ["All bookings", "More than 60 days", '=COUNTIFS(rng_lead,">60")', '=COUNTIFS(rng_lead,">60",rng_cancel,1)',
     f"=D{b1 + 2}/C{b1 + 2}", f"=E{b1 + 2}/E{b1 + 3}", f"=F{b1 + 2}-1"],
    ["All bookings", "60 days or less", '=COUNTIFS(rng_lead,"<=60")', '=COUNTIFS(rng_lead,"<=60",rng_cancel,1)',
     f"=D{b1 + 3}/C{b1 + 3}", "reference", "reference"],
]
table_block(q, TB_H, 1, hdrB, rowsB, ["@", "@", FMT_N, FMT_N, FMT_P, FMT_RR, "0%"], hdr_height=48)
for r in (b1 + 1, b1 + 3):
    for col in "FG":
        q[f"{col}{r}"].alignment = A_CTR
        q[f"{col}{r}"].font = Font(name=FONT, size=10, italic=True, color="7F7F7F")
TB_END = b1 + 3
note(q, f"A{TB_END + 1}", "Relative risk = cancellation rate for bookings made more than 60 days ahead ÷ rate for 60 days "
                          "or less; '% more likely' = relative risk − 1. D-EDGE Hospitality Solutions (2019) reports "
                          "65% across 680 European hotels. Source: Clean_Data.")

# Table C
TC_T = TB_END + 3
ttitle(q, f"A{TC_T}", "Table C: Robustness and cross-checks")
put(q, f"A{TC_T + 1}", "C1. Bookings per band computed two ways", font=F_BOLD)
C1_H = TC_T + 2
c10 = C1_H + 1
rowsC1 = []
for i, (lab, lo, hi) in enumerate(BANDS):
    r = c10 + i
    rowsC1.append([f"=A{TA0 + i}", f"=D{TA0 + i}", f"=COUNTIF(rng_band,A{r})", f"=C{r}-B{r}"])
C1_TOT = c10 + len(BANDS)
rowsC1.append(["Total", f"=SUM(B{c10}:B{C1_TOT - 1})", f"=SUM(C{c10}:C{C1_TOT - 1})", f"=C{C1_TOT}-B{C1_TOT}"])
table_block(q, C1_H, 1, ["Lead time band (days)", "Table A bookings (lead_time bounds)",
                         "COUNTIF on lead_band", "Difference (must be 0)"],
            rowsC1, ["@", FMT_N, FMT_N, FMT_N], total_rows=(len(BANDS),), hdr_height=48)

put(q, f"A{C1_TOT + 2}", "C2. Bookings made 31–365 days ahead", font=F_BOLD)
C2_H = C1_TOT + 3
C2 = C2_H + 1
table_block(q, C2_H, 1, ["Group", "Bookings", "Share of all bookings", "Cancelled", "Share of all cancellations"],
            [["31–365 days", '=COUNTIFS(rng_lead,">=31",rng_lead,"<=365")', f"=B{C2}/$D${TA_TOT}",
              '=COUNTIFS(rng_lead,">=31",rng_lead,"<=365",rng_cancel,1)', f"=D{C2}/$E${TA_TOT}"]],
            ["@", FMT_N, FMT_P, FMT_N, FMT_P], hdr_height=33)

put(q, f"A{C2 + 2}", "C3. Excluding duplicate rows (dup_flag = 0)", font=F_BOLD)
C3_H = C2 + 3
c30 = C3_H + 1
rowsC3 = []
for i, (lab, lo, hi) in enumerate(BANDS):
    r, ra = c30 + i, TA0 + i
    band = f'rng_lead,">="&$B${ra},rng_lead,"<="&$C${ra}'
    rowsC3.append([f"=A{ra}", f"=COUNTIFS({band},rng_dup,0)", f"=COUNTIFS({band},rng_dup,0,rng_cancel,1)",
                   f"=C{r}/B{r}", f"=F{ra}", f"=E{r}-D{r}"])
C3_TOT = c30 + len(BANDS)
rowsC3.append(["Total", f"=SUM(B{c30}:B{C3_TOT - 1})", f"=SUM(C{c30}:C{C3_TOT - 1})", f"=C{C3_TOT}/B{C3_TOT}",
               f"=F{TA_TOT}", f"=E{C3_TOT}-D{C3_TOT}"])
table_block(q, C3_H, 1, ["Lead time band (days)", "Bookings (dup_flag = 0)", "Cancelled", "Cancellation rate",
                         "All-bookings rate (Table A)", "Gap (all minus de-duplicated)"],
            rowsC3, ["@", FMT_N, FMT_N, FMT_P, FMT_P, FMT_P], total_rows=(len(BANDS),), hdr_height=48)
note(q, f"A{C3_TOT + 1}", "C1 checks the formula band column against direct lead_time bounds. C3 repeats the analysis "
                          "without rows that duplicate an earlier row on all 28 fields; the gradient still rises. "
                          "Source: Clean_Data.")

# Table D
TD_T = C3_TOT + 3
TD_H = TD_T + 1
d0 = TD_H + 1
ttitle(q, f"A{TD_T}", "Table D: Chart data (feeds Figure 2)")
rowsD = []
for i in range(len(BANDS)):
    r, ra = d0 + i, TA0 + i
    label = (f'=A{ra}&CHAR(10)&TEXT(D{ra},"#,##0")&" bookings"&CHAR(10)&TEXT(E{ra},"#,##0")&" cancelled"')
    if i == len(BANDS) - 1:  # Over 365: add the non-refundable share that explains the no-deposit drop
        label += f'&CHAR(10)&TEXT(N{ra},"0%")&" non-refundable"'
    rowsD.append([label, f"=F{ra}", f"=I{ra}", f"=AND(B{r}=F{ra},C{r}=I{ra})"])
TD_END = d0 + len(BANDS) - 1
SER_NAMES = ["Cancellation rate, all bookings", "Cancellation rate, no-deposit bookings"]  # = legend entries
table_block(q, TD_H, 1, ["Category label (axis)"] + SER_NAMES + ["Check: equals Table A?"],
            rowsD, ["@", FMT_P, FMT_P, None], hdr_height=48)
for i, r in enumerate(range(d0, TD_END + 1)):
    q[f"A{r}"].alignment = Alignment(wrap_text=True, vertical="center")
    q.row_dimensions[r].height = 62 if i == len(BANDS) - 1 else 47
    q[f"D{r}"].alignment = A_CTR
bool_rules(q, f"D{d0}:D{TD_END}")
note(q, f"A{TD_END + 1}", "Each bar is the share of bookings in that band that were cancelled (is_canceled = Yes, "
                          "including no-shows). Labels show each band's bookings and cancellations. "
                          "Source: Table A above.")
presentation(q, "70AD47", area=f"A1:P{TD_END + 1}")
q.row_breaks.append(Break(id=TC_T - 1))  # Table C starts a new printed page
q.freeze_panes = "B4"

# ---------------------------------------------------------------- 5. Table1_DataQuality
t = ws_t1
title(t, "A1", "Table 1 (working): data quality checks with live formula counts")
put(t, "A2", "Each count is a live formula on Raw_Data or Clean_Data. 'Expected' values are typed constants from an "
             "independent Python check (pandas and a plain openpyxl row loop).", font=F_SMALL)
ttitle(t, "A3", "Part 1: Formula checks")
T1_H = 4
seg, chan, dep, stat, canc = raw("market_segment"), raw("distribution_channel"), raw("deposit_type"), \
    raw("reservation_status"), raw("is_canceled")
adr = raw("Average_Daily_Rate")
checks1 = [  # key, label, field, formula, expected, fmt
    ("rows", "Rows in file (bookings)", "booking_order_id", f"=COUNTA({raw('booking_order_id')})", 119390, FMT_N),
    ("city", "City Hotel bookings", "hotel_name", f'=COUNTIF({raw("hotel_name")},"City Hotel")', 79330, FMT_N),
    ("resort", "Resort Hotel bookings", "hotel_name", f'=COUNTIF({raw("hotel_name")},"Resort Hotel")', 40060, FMT_N),
    ("yesno", "is_canceled coded as text Yes/No", "is_canceled",
     f'=COUNTIF({canc},"Yes")+COUNTIF({canc},"No")', 119390, FMT_N),
    ("rep_yesno", "is_repeated_guest coded as text Yes/No", "is_repeated_guest",
     f'=COUNTIF({raw("is_repeated_guest")},"Yes")+COUNTIF({raw("is_repeated_guest")},"No")', 119390, FMT_N),
    ("canc", "is_canceled = Yes", "is_canceled", f'=COUNTIF({canc},"Yes")', 44224, FMT_N),
    ("st_canc", "…of which reservation_status = Canceled", "reservation_status", f'=COUNTIF({stat},"Canceled")', 43017, FMT_N),
    ("st_ns", "…of which reservation_status = No-Show", "reservation_status", f'=COUNTIF({stat},"No-Show")', 1207, FMT_N),
    ("st_co", "reservation_status = Check-Out", "reservation_status", f'=COUNTIF({stat},"Check-Out")', 75166, FMT_N),
    ("mismatch", "is_canceled vs reservation_status mismatches", "is_canceled, reservation_status",
     f'=SUMPRODUCT(({canc}="Yes")*({stat}="Check-Out"))+SUMPRODUCT(({canc}="No")*({stat}<>"Check-Out"))', 0, FMT_N),
    ("lt_blank", "lead_time blanks", "lead_time", f"=COUNTBLANK({raw('lead_time')})", 0, FMT_N),
    ("lt_neg", "lead_time negative values", "lead_time", f'=COUNTIF({raw("lead_time")},"<0")', 0, FMT_N),
    ("lt_min", "lead_time minimum (days)", "lead_time", f"=MIN({raw('lead_time')})", 0, FMT_N),
    ("lt_max", "lead_time maximum (days)", "lead_time", f"=MAX({raw('lead_time')})", 737, FMT_N),
    ("lt_med", "lead_time median (days)", "lead_time", f"=MEDIAN({raw('lead_time')})", 69, "0.0"),
    ("lt_mean", "lead_time mean (days)", "lead_time", f"=AVERAGE({raw('lead_time')})", None, "0.0000"),
    ("dups", "Duplicate rows (dup_flag = 1)", "all 28 fields", "=SUM(rng_dup)", 32239, FMT_N),
    ("dup_go", "Duplicates that are Groups or Offline TA/TO", "market_segment",
     '=COUNTIFS(rng_dup,1,rng_segment,"Groups")+COUNTIFS(rng_dup,1,rng_segment,"Offline TA/TO")', 25313, FMT_N),
    ("dup_go_sh", "…as a share of all duplicates", "market_segment", "={dup_go}/{dups}", None, FMT_P),
    ("go", "Groups + Offline TA/TO bookings", "market_segment",
     f'=COUNTIF({seg},"Groups")+COUNTIF({seg},"Offline TA/TO")', 44030, FMT_N),
    ("go_sh", "…as a share of all bookings", "market_segment", "={go}/{rows}", None, FMT_P),
    ("country", "Missing country", "country", f"=COUNTBLANK({raw('country')})", 488, FMT_N),
    ("seg_u", "market_segment = Undefined", "market_segment", f'=COUNTIF({seg},"Undefined")', 2, FMT_N),
    ("chan_u", "distribution_channel = Undefined", "distribution_channel", f'=COUNTIF({chan},"Undefined")', 5, FMT_N),
    ("undef", "Rows with segment or channel = Undefined", "market_segment, distribution_channel",
     f'=SUMPRODUCT(--((({seg}="Undefined")+({chan}="Undefined"))>0))', 5, FMT_N),
    ("agent", "Blank agent", "agent", f"=COUNTBLANK({raw('agent')})", 16340, FMT_N),
    ("company", "Blank company", "company", f"=COUNTBLANK({raw('company')})", 112593, FMT_N),
    ("ac_num", "agent or company stored as numbers (non-blank)", "agent, company",
     f"=SUMPRODUCT(--((ISNUMBER({raw('agent')})+ISNUMBER({raw('company')}))>0))", None, FMT_N),
    ("dates", "Invalid arrival dates (arrival_date_valid = 0)", "arrival_date", "=COUNTIF(rng_valid,0)", 96, FMT_N),
    ("children", "Missing children", "children", f"=COUNTBLANK({raw('children')})", 4, FMT_N),
    ("zero", "Zero guests (guests_total = 0)", "adults, children, babies", "=COUNTIF(rng_guests,0)", 180, FMT_N),
    ("adr_min", "Average_Daily_Rate minimum", "Average_Daily_Rate", f"=MIN({adr})", -6.38, "#,##0.00"),
    ("adr_max", "Average_Daily_Rate maximum", "Average_Daily_Rate", f"=MAX({adr})", 5400, "#,##0.00"),
    ("adr_neg", "Bookings with negative Average_Daily_Rate", "Average_Daily_Rate", f'=COUNTIF({adr},"<0")', 1, FMT_N),
    ("adr_top", "Bookings at the maximum Average_Daily_Rate", "Average_Daily_Rate",
     f"=COUNTIF({adr},MAX({adr}))", 1, FMT_N),
    ("odd", "Distinct bookings with missing children, zero guests or an extreme rate", "children, guests_total, ADR",
     f"=SUMPRODUCT(--((({raw('children')}=\"\")+(rng_guests=0)+({adr}<0)+({adr}=MAX({adr})))>0))", None, FMT_N),
    ("nodep", "deposit_type = No Deposit", "deposit_type", f'=COUNTIF({dep},"No Deposit")', 104641, FMT_N),
    ("nonref", "deposit_type = Non Refund", "deposit_type", f'=COUNTIF({dep},"Non Refund")', 14587, FMT_N),
    ("refund", "deposit_type = Refundable", "deposit_type", f'=COUNTIF({dep},"Refundable")', 162, FMT_N),
]
# expected values not listed in the brief come from the independent Python check (verify.py repeats them)
EXTRA_EXPECTED = {"lt_mean": round(104.0114, 4), "dup_go_sh": round(25313 / 32239, 4), "go_sh": round(44030 / 119390, 4),
                  "ac_num": 109630, "odd": 186}
T1ROW = {}
r = T1_H + 1
for key, *_ in checks1:
    T1ROW[key] = r
    r += 1
rows1 = []
for i, (key, label, field, formula, expected, fmt) in enumerate(checks1):
    rr = T1_H + 1 + i
    formula = formula.format(**{k: f"D{v}" for k, v in T1ROW.items()})
    if expected is None:
        expected = EXTRA_EXPECTED[key]
    rows1.append([i + 1, label, field, formula, expected, f'=IF(ROUND(D{rr},4)=ROUND(E{rr},4),"{OK}","{BAD}")'])
T1_END = T1_H + len(rows1)
table_block(t, T1_H, 1, ["#", "Check", "Field(s)", "Bookings affected (live formula)",
                         "Expected, from independent Python check", "Match?"],
            rows1, ["0", "@", "@", None, None, "@"], widths=[14, 50, 22, 16, 48, 11], hdr_height=48,
            wrap_cols=(1, 2))
for i, (key, label, field, formula, expected, fmt) in enumerate(checks1):
    rr = T1_H + 1 + i
    t[f"D{rr}"].number_format = fmt
    t[f"E{rr}"].number_format = fmt
    t[f"E{rr}"].font = Font(name=FONT, size=11, color="0000FF")
    t[f"F{rr}"].alignment = A_CTR
    t[f"A{rr}"].alignment = A_CTR
match_rules(t, f"F{T1_H + 1}:F{T1_END}")
note(t, f"A{T1_END + 1}", "Blue figures are typed constants (expected values from verify.py); black figures are formulas. "
                          "Match? compares both after rounding to 4 decimal places.")

# Presentation version (copy-ready for Word), mirroring report Table 1
P_T = T1_END + 3
ttitle(t, f"A{P_T}", "Part 2: Table 1 as it appears in the report (copy-ready)")
put(t, f"A{P_T + 1}", "Table 1: Data quality checks and treatment", font=F_BOLD)
D = {k: f"D{v}" for k, v in T1ROW.items()}
QS = "Q2_Summary"
pres = [
    ["Consistency", "is_canceled and is_repeated_guest are coded Yes/No, not 0/1 as the data dictionary states",
     f"={D['yesno']}", "Yes", "Recoded to 1/0, so the average of is_canceled gives the cancellation rate"],
    ["Consistency",
     f'="is_canceled = Yes covers "&TEXT({D["st_canc"]},"#,##0")&" cancellations and "&TEXT({D["st_ns"]},"#,##0")'
     f'&" no-shows, and it "&IF({D["mismatch"]}=0,"agrees with reservation_status for every booking",'
     f'"disagrees with reservation_status for "&TEXT({D["mismatch"]},"#,##0")&" bookings")',
     f"={D['canc']}", "Yes",
     'Kept as defined: "cancellation" in this report includes no-shows unless stated'],
    ["Validity",
     f'="lead_time has no blanks or negative values ("&{D["lt_min"]}&" to "&{D["lt_max"]}&" days) but is '
     f'right-skewed (median "&TEXT({D["lt_med"]},"0")&", mean "&TEXT({D["lt_mean"]},"0")&" days)"',
     f"={D['lt_blank']}+{D['lt_neg']}", "Yes (Q2)",
     "Grouped into six bands rather than averaged, since a single average hides the shape of skewed data "
     "(Vidgen, Kirshner & Tan 2019, pp. 54, 85); António, Almeida and Nunes (2017, p. 31) report the same skew "
     "at four Algarve resorts"],
    ["Consistency", "Rows identical on all 28 variables but with different booking IDs", f"={D['dups']}", "Yes",
     f'="Kept: "&TEXT({D["dup_go_sh"]},"0.0%")&" are Groups or Offline TA/TO bookings ("&TEXT({D["go_sh"]},"0.0%")'
     f'&" of all bookings), so they most likely record separate rooms in multi-room bookings. The Q2 pattern '
     f'holds without them ("&TEXT({QS}!D{c30},"0.0%")&" at 0–7 days rising to "&TEXT({QS}!D{c30 + 4},"0.0%")'
     f'&" at 181–365 days)"'],
    ["Completeness", "Missing country", f"={D['country']}", "Yes (Q1)",
     'Labelled "Unknown" and counted in all-guest totals'],
    ["Completeness", f'="market_segment ("&{D["seg_u"]}&") or distribution_channel ("&{D["chan_u"]}'
                     f'&") recorded as ""Undefined"""', f"={D['undef']}", "Yes (Q3, Q4)",
     'Kept as "Undefined"; too few to move any rate'],
    ["Completeness", "Blank agent and company", f'=TEXT({D["agent"]},"#,##0")&" / "&TEXT({D["company"]},"#,##0")',
     "No", 'Recoded "None": a blank means no agent or company was involved (Antonio, de Almeida & Nunes 2019c), '
           "and an optional field left blank is still complete (Vidgen, Kirshner & Tan 2019, p. 57)"],
    ["Validity", "agent and company are stored as numbers but are ID codes", f"={D['rows']}", "No",
     "Treated as dimensions"],
    ["Validity", 'arrival_date holds full dates, although the dictionary says "Year of arrival". The original '
                 "2015–2017 dates were moved forward seven years, turning 29 February 2016 into the impossible "
                 'text date "29 February 2023"', f"={D['dates']}", "No",
     "Flagged; excluded from any date-based analysis"],
    ["Validity", f'="Missing children ("&{D["children"]}&"), no guests recorded ("&{D["zero"]}'
                 f'&"), Average_Daily_Rate of "&TEXT({D["adr_min"]},"0.00")&" ("&{D["adr_neg"]}&") and "'
                 f'&TEXT({D["adr_max"]},"#,##0")&" ("&{D["adr_top"]}&")"',
     f"={D['children']}+{D['zero']}+{D['adr_neg']}+{D['adr_top']}", "No",
     "Flagged and kept, since these fields enter no business question"],
]
P_H = P_T + 2
P_END = table_block(t, P_H, 1, ["Quality dimension", "Check and issue found", "Bookings affected",
                                "Used in Q1–Q5?", "Treatment"],
                    pres, ["@", "@", FMT_N, "@", "@"], hdr_height=33, wrap_cols=(1, 4))
for i, rr in enumerate(range(P_H + 1, P_END + 1)):
    t[f"C{rr}"].alignment = Alignment(horizontal="right", vertical="top")
    t[f"A{rr}"].alignment = Alignment(vertical="top")
    t[f"D{rr}"].alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
    # height from the longest text (formula text overstates a little, which only adds white space)
    lines = max(-(-shown_len(pres[i][1]) // 50), -(-shown_len(pres[i][4]) // 48), 1)
    t.row_dimensions[rr].height = 15 * min(lines, 7) + 4
note(t, f"A{P_END + 1}", "Source: Group analysis of the supplied hotel bookings dataset (n = 119,390); quality "
                         "dimensions from Vidgen, Kirshner and Tan (2019, pp. 57–58).")
t[f"A{P_END + 1}"].value = (f'="Source: Group analysis of the supplied hotel bookings dataset (n = "&TEXT({D["rows"]},"#,##0")'
                            f'&"); quality dimensions from Vidgen, Kirshner and Tan (2019, pp. 57–58)."')
put(t, f"A{P_END + 2}", "Note: 'Bookings affected' for the last row adds the four counts; the distinct number of "
                        "bookings involved is shown in Part 1 (check 'Distinct bookings with missing children, zero "
                        "guests or an extreme rate').", font=F_NOTE)
presentation(t, "70AD47", area=f"A1:F{P_END + 2}")
t.row_breaks.append(Break(id=P_T - 1))  # the copy-ready table starts a new printed page
t.freeze_panes = "A4"

# ---------------------------------------------------------------- 6. Figure2 (native clustered column chart)
f = ws_f
put(f, "A1", "Figure 2: Cancellation rate by booking lead time, all bookings and no-deposit bookings",
    font=Font(name=FONT, size=12, bold=True))
put(f, "A2", "(For Word: the caption goes above the pasted chart; the note and source go below it.)",
    font=Font(name=FONT, size=10, italic=True, color="595959"))

BLUE, GREEN, AXIS_TXT, AXIS_LINE = "0072B2", "009E73", "595959", "BFBFBF"


def cp(size=900, color="404040", bold=False):
    return CharacterProperties(sz=size, b=bold, solidFill=color, latin=DFont(typeface=FONT),
                               cs=DFont(typeface=FONT))


def rich(size=900, color="404040", bold=False, rot=None):
    return RichText(bodyPr=RichTextProperties(rot=rot, vert="horz") if rot is not None else RichTextProperties(),
                    p=[Paragraph(pPr=ParagraphProperties(defRPr=cp(size, color, bold)), r=[],
                                 endParaRPr=cp(size, color, bold))])


def axis_title(text, size=1000, color="262626", bold=True):
    para = Paragraph(pPr=ParagraphProperties(defRPr=cp(size, color, bold)),
                     r=[RegularTextRun(rPr=cp(size, color, bold), t=text)])
    return Title(tx=Text(rich=RichText(bodyPr=RichTextProperties(), p=[para])), overlay=False)


def figure2_write(chart):
    """Serialise like BarChart, then add chart-wide Calibri text, autoTitleDeleted and sourceLinked=0 formats."""
    tree = BarChart._write(chart)  # tags are bare here; openpyxl adds the chart namespace when serialising
    from openpyxl.xml.functions import Element
    tree.find("chart").insert(0, Element("autoTitleDeleted", val="1"))
    txpr = rich(900).to_tree(tagname="txPr")
    children = list(tree)
    sppr_pos = [i for i, el in enumerate(children) if el.tag == "spPr"]
    tree.insert(sppr_pos[0] + 1 if sppr_pos else len(children), txpr)
    for nf in tree.iter("numFmt"):
        nf.set("sourceLinked", "0")
    return tree


def build_figure2(data_ws, header_row, first_row, last_row, cat_col, all_col, nd_col, target_ws, anchor):
    """Figure 2 as in figure2_final_reference.py: two clustered columns, value labels, hidden 0-80% value axis."""
    ch = BarChart()
    ch._write = lambda: figure2_write(ch)
    ch.type = "col"
    ch.grouping = "clustered"
    ch.varyColors = False
    sheet = f"'{data_ws.title}'"
    cat_ref = f"{sheet}!${get_column_letter(cat_col)}${first_row}:${get_column_letter(cat_col)}${last_row}"
    for col_idx, colour in ((all_col, BLUE), (nd_col, GREEN)):
        col = get_column_letter(col_idx)
        s = Series(Reference(data_ws, min_col=col_idx, min_row=first_row, max_row=last_row), title=None)
        s.tx = SeriesLabel(strRef=StrRef(f"{sheet}!${col}${header_row}"))  # legend text = header cell
        s.cat = AxDataSource(strRef=StrRef(cat_ref))                       # text categories (multi-line labels)
        s.graphicalProperties = GraphicalProperties(solidFill=colour, ln=LineProperties(noFill=True))
        s.dLbls = DataLabelList(numFmt="0.0%", dLblPos="outEnd", showVal=True, showSerName=False,
                                showCatName=False, showLegendKey=False, showPercent=False, showBubbleSize=False,
                                txPr=rich(800, colour, True))
        ch.series.append(s)
    ch.gapWidth = 40
    ch.overlap = -12
    # value axis: fixed 0-80% scale (honest zero baseline), hidden because every bar carries its own label
    ch.y_axis.scaling.min = 0
    ch.y_axis.scaling.max = 0.8
    ch.y_axis.majorUnit = 0.2
    ch.y_axis.number_format = "0%"
    ch.y_axis.majorGridlines = None
    ch.y_axis.minorGridlines = None
    ch.y_axis.delete = True
    # category axis: band / bookings / cancelled (Over 365 adds the non-refundable share)
    ch.x_axis.title = axis_title("Booking lead time (days before arrival)")
    ch.x_axis.txPr = rich(800, AXIS_TXT, rot=0)  # rot=0 stops Excel tilting the labels
    ch.x_axis.spPr = GraphicalProperties(ln=LineProperties(solidFill=AXIS_LINE, w=9525))
    ch.x_axis.delete = False
    ch.x_axis.majorTickMark = "none"
    ch.x_axis.tickLblSkip = 1  # never skip alternate labels
    ch.x_axis.tickMarkSkip = 1
    ch.legend.position = "t"
    ch.legend.overlay = False
    ch.legend.txPr = rich(900, "404040")
    ch.title = None
    ch.roundedCorners = False
    ch.graphical_properties = GraphicalProperties(solidFill="FFFFFF", ln=LineProperties(noFill=True))
    ch.plot_area.graphicalProperties = GraphicalProperties(noFill=True, ln=LineProperties(noFill=True))
    ch.width, ch.height = 15.9, 9.5  # A4 text width at 2.54 cm margins, so it pastes into Word at 100%
    target_ws.add_chart(ch, anchor)
    return ch


build_figure2(q, header_row=TD_H, first_row=d0, last_row=TD_END, cat_col=1, all_col=2, nd_col=3,
              target_ws=f, anchor="A4")

NOTE_ROW = 24
f[f"A{NOTE_ROW}"].value = (f'="Note: Each bar is the share of bookings in that lead-time band that were cancelled, '
                           f'including no-shows. Labels under each band give its bookings and cancellations. Beyond '
                           f'365 days, "&TEXT({QS}!N{TA_TOT - 1},"0.0%")&" of bookings are non-refundable, so most of '
                           f'that band sits outside the no-deposit bar."')
f[f"A{NOTE_ROW + 1}"].value = (f'="Source: Group analysis of the supplied hotel bookings dataset (n = "'
                               f'&TEXT({QS}!D{TA_TOT},"#,##0")&")."')
for col in range(1, 11):  # ten columns ≈ the 15.9 cm chart width
    f.column_dimensions[get_column_letter(col)].width = 8.0
for rr, height in ((1, 32), (2, 15), (NOTE_ROW, 38), (NOTE_ROW + 1, 14)):
    f.merge_cells(f"A{rr}:J{rr}")
    f[f"A{rr}"].alignment = Alignment(wrap_text=True, vertical="top", horizontal="left")
    f.row_dimensions[rr].height = height
for rr in (NOTE_ROW, NOTE_ROW + 1):
    f[f"A{rr}"].font = Font(name=FONT, size=9, color="404040")
presentation(f, "ED7D31", landscape=False, area="A1:J27")
f.page_setup.fitToHeight = 1  # fit to one page

# ---------------------------------------------------------------- 7. Report_Figures
rf = ws_rf
title(rf, "A1", "Report figures: every number quoted in Section 6.2 / Table 1, traced to its live cell")
put(rf, "A2", "Match compares ROUND(live value, 3) with the expected value (ROUND(…, 2) for the 62% row).", font=F_SMALL)
RF_H = 4
T1 = "Table1_DataQuality"
rf_items = [
    ("0–7 day cancellation rate, all bookings", "6.2", f"={QS}!F{TA0}", 0.096, 3),
    ("181–365 day cancellation rate, all bookings", "6.2", f"={QS}!F{TA0 + 4}", 0.555, 3),
    ("Over-365 cancellation rate, all bookings", "6.2", f"={QS}!F{TA0 + 5}", 0.677, 3),
    ("31–365 day share of bookings", "6.2", f"={QS}!C{C2}", 0.649, 3),
    ("31–365 day share of cancellations", "6.2", f"={QS}!E{C2}", 0.789, 3),
    ("Over-365 Non Refund share", "6.2, Figure 2 note", f"={QS}!N{TA0 + 5}", 0.592, 3),
    ("Over-365 no-deposit cancellation rate", "6.2", f"={QS}!I{TA0 + 5}", 0.206, 3),
    ("No-deposit Online TA, 0–7 day rate", "6.2", f"={QS}!L{TA0}", 0.092, 3),
    ("No-deposit Online TA, 181–365 day rate", "6.2", f"={QS}!L{TA0 + 4}", 0.534, 3),
    ("No-deposit >60 vs ≤60 days, % more likely to cancel", "6.2", f"={QS}!G{b1}", 0.62, 2),
    ("De-duplicated 0–7 day rate", "Table 1", f"={QS}!D{c30}", 0.084, 3),
    ("De-duplicated 181–365 day rate", "Table 1", f"={QS}!D{c30 + 4}", 0.393, 3),
    ("Duplicates that are Groups or Offline TA/TO", "Table 1", f"={T1}!D{T1ROW['dup_go_sh']}", 0.785, 3),
    ("Groups + Offline TA/TO share of all bookings", "Table 1", f"={T1}!D{T1ROW['go_sh']}", 0.369, 3),
]
rows_rf = []
for i, (lab, sec, ref, exp, dp) in enumerate(rf_items):
    rr = RF_H + 1 + i
    rows_rf.append([lab, sec, ref, exp, f'=IF(ROUND(C{rr},{dp})=D{rr},"{OK}","{BAD}")'])
RF_END = table_block(rf, RF_H, 1, ["Report text", "Where quoted", "Live value (formula link)", "Expected (report)",
                                   "Match"], rows_rf, ["@", "@", FMT_P, FMT_P, "@"], widths=[52, 18, 16, 16, 11],
                     hdr_height=33)
for rr in range(RF_H + 1, RF_END + 1):
    rf[f"D{rr}"].font = Font(name=FONT, size=11, color="0000FF")
    rf[f"E{rr}"].alignment = A_CTR
rf[f"C{RF_H + 10}"].number_format = "0.0%"
rf[f"D{RF_H + 10}"].number_format = "0%"
match_rules(rf, f"E{RF_H + 1}:E{RF_END}")
note(rf, f"A{RF_END + 1}", "Blue = expected value typed from the report text; black = live formula. Source: Q2_Summary and "
                           "Table1_DataQuality.")
presentation(rf, "C00000", area=f"A1:E{RF_END + 1}")
rf.freeze_panes = f"A{RF_H + 1}"

# ---------------------------------------------------------------- 8. Checks
c = ws_ch
title(c, "A1", "Self-checks dashboard")
put(c, "A2", "Every result is a live TRUE/FALSE formula. The workbook is consistent only when the banner reads "
             "ALL CHECKS PASS.", font=F_SMALL)
CH_H = 5
ck = []
ck.append(("Sum of band bookings = COUNTA(booking IDs)", f"={QS}!D{TA_TOT}", "=COUNTA(rng_id)", "=C{r}=D{r}"))
ck.append(("Sum of band bookings = 119,390", f"={QS}!D{TA_TOT}", 119390, "=C{r}=D{r}"))
ck.append(("Sum of band cancellations = SUM(cancel_flag)", f"={QS}!E{TA_TOT}", "=SUM(rng_cancel)", "=C{r}=D{r}"))
ck.append(("SUM(cancel_flag) = COUNTIF(is_canceled, \"Yes\")", "=SUM(rng_cancel)", '=COUNTIF(rng_iscanceled,"Yes")',
           "=C{r}=D{r}"))
for i, (lab, lo, hi) in enumerate(BANDS):
    ck.append((f"Band {lab}: Table A bookings = COUNTIF(lead_band) (Table C1)", f"={QS}!D{TA0 + i}",
               f"={QS}!C{c10 + i}", "=C{r}=D{r}"))
ck.append(("No-deposit total = COUNTIF(deposit_type, \"No Deposit\")", f"={QS}!G{TA_TOT}",
           '=COUNTIF(rng_deposit,"No Deposit")', "=C{r}=D{r}"))
ck.append(("Non Refund total = COUNTIF(deposit_type, \"Non Refund\")", f"={QS}!M{TA_TOT}",
           '=COUNTIF(rng_deposit,"Non Refund")', "=C{r}=D{r}"))
rate_ranges = (f"{QS}!F{TA0}:F{TA_TOT},{QS}!I{TA0}:I{TA_TOT},{QS}!L{TA0}:L{TA_TOT - 1},{QS}!N{TA0}:N{TA_TOT},"
               f"{QS}!O{TA0}:P{TA_TOT - 1},{QS}!E{b1}:E{b1 + 3},{QS}!D{c30}:D{C3_TOT},{QS}!C{C2},{QS}!E{C2}")
ck.append(("Every rate and share lies between 0 and 1 (minimum)", f"=MIN({rate_ranges})", 0, "=C{r}>=D{r}"))
ck.append(("Every rate and share lies between 0 and 1 (maximum)", f"=MAX({rate_ranges})", 1, "=C{r}<=D{r}"))
ck.append(("Shares of bookings sum to 100% (within 1E-9)", f"=SUM({QS}!O{TA0}:O{TA_TOT - 1})", 1,
           "=ABS(C{r}-D{r})<1E-9"))
ck.append(("Shares of cancellations sum to 100% (within 1E-9)", f"=SUM({QS}!P{TA0}:P{TA_TOT - 1})", 1,
           "=ABS(C{r}-D{r})<1E-9"))
for i in range(len(BANDS) - 1):
    ck.append((f"All-bookings rate rises: {BANDS[i][0]} < {BANDS[i + 1][0]}", f"={QS}!F{TA0 + i}",
               f"={QS}!F{TA0 + i + 1}", "=C{r}<D{r}"))
ck.append(("De-duplicated bookings = COUNTA − SUM(dup_flag)", f"={QS}!B{C3_TOT}", "=COUNTA(rng_id)-SUM(rng_dup)",
           "=C{r}=D{r}"))
ck.append(("De-duplicated bookings = 119,390 − 32,239 = 87,151", f"={QS}!B{C3_TOT}", 87151, "=C{r}=D{r}"))
ck.append(("Table B no-deposit >60 + ≤60 = no-deposit total", f"={QS}!C{b1}+{QS}!C{b1 + 1}", f"={QS}!G{TA_TOT}",
           "=C{r}=D{r}"))
ck.append(("Table B all bookings >60 = SUM(lead_over_60)", f"={QS}!C{b1 + 2}", "=SUM(rng_over60)", "=C{r}=D{r}"))
ck.append(("Table B all >60 + ≤60 = all bookings", f"={QS}!C{b1 + 2}+{QS}!C{b1 + 3}", f"={QS}!D{TA_TOT}", "=C{r}=D{r}"))
ck.append(("31–365 combined (C2) = Table A bands 31–90 to 181–365", f"={QS}!B{C2}", f"=SUM({QS}!D{TA0 + 2}:D{TA0 + 4})",
           "=C{r}=D{r}"))
ck.append(("Table D series equal Table A in all six bands", f"=COUNTIF({QS}!D{d0}:D{TD_END},TRUE)", 6,
           "=C{r}=D{r}"))
ck.append(("Every Figure 2 axis label shows that band's bookings",
           f'=SUMPRODUCT(--ISNUMBER(SEARCH(TEXT({QS}!D{TA0}:D{TA_TOT - 1},"#,##0")&" bookings",'
           f"{QS}!A{d0}:A{TD_END})))", 6, "=C{r}=D{r}"))
ck.append(("Every Figure 2 axis label shows that band's cancellations",
           f'=SUMPRODUCT(--ISNUMBER(SEARCH(TEXT({QS}!E{TA0}:E{TA_TOT - 1},"#,##0")&" cancelled",'
           f"{QS}!A{d0}:A{TD_END})))", 6, "=C{r}=D{r}"))
ck.append(("Over 365 axis label shows the non-refundable share",
           f'=ISNUMBER(SEARCH(TEXT({QS}!N{TA_TOT - 1},"0%")&" non-refundable",{QS}!A{TD_END}))', "=TRUE",
           "=C{r}=D{r}"))
ck.append(("Chart scale fits every value (max rate < 80%)", f"=MAX({QS}!B{d0}:C{TD_END})", 0.8, "=C{r}<D{r}"))
ck.append(("cancel_flag has no errors (pivot-ready)", "=SUMPRODUCT(--ISERROR(rng_cancel))", 0, "=C{r}=D{r}"))
ck.append(("Raw_Data rows = Clean_Data rows", f"=COUNTA(Raw_Data!$A$2:$A${LAST})", "=COUNTA(rng_id)", "=C{r}=D{r}"))
ck.append(("Invalid dates flagged = text cells in Raw_Data arrival_date", "=COUNTIF(rng_valid,0)",
           f"=SUMPRODUCT(--ISTEXT({raw('arrival_date')}))", "=C{r}=D{r}"))
ck.append(("Table1_DataQuality: every Match? cell says OK", f'=COUNTIF({T1}!F{T1_H + 1}:F{T1_END},"{OK}")',
           f"=ROWS({T1}!F{T1_H + 1}:F{T1_END})", "=C{r}=D{r}"))
ck.append(("Report_Figures: every Match cell says OK", f'=COUNTIF(Report_Figures!E{RF_H + 1}:E{RF_END},"{OK}")',
           f"=ROWS(Report_Figures!E{RF_H + 1}:E{RF_END})", "=C{r}=D{r}"))
rows_ck = []
for i, (lab, left, right, test) in enumerate(ck):
    rr = CH_H + 1 + i
    rows_ck.append([i + 1, lab, left, right, test.format(r=rr)])
CH_END = table_block(c, CH_H, 1, ["#", "Check", "Value", "Compared with", "Result"], rows_ck,
                     ["0", "@", "#,##0.######", "#,##0.######", "@"], widths=[5, 62, 16, 16, 12], hdr_height=22)
for rr in range(CH_H + 1, CH_END + 1):
    c[f"A{rr}"].alignment = A_CTR
    c[f"E{rr}"].alignment = A_CTR
bool_rules(c, f"E{CH_H + 1}:E{CH_END}")
RES = f"E{CH_H + 1}:E{CH_END}"
put(c, "A3", f'=IF(AND(COUNTIF({RES},FALSE)=0,COUNTIF({RES},TRUE)=ROWS({RES})),"ALL CHECKS PASS","CHECK FAILED")',
    font=Font(name=FONT, size=20, bold=True), align=Alignment(horizontal="left", vertical="center"))
c.row_dimensions[3].height = 32
c.conditional_formatting.add("A3", FormulaRule(formula=['A3="ALL CHECKS PASS"'], font=Font(name=FONT, size=20, bold=True, color="006100")))
c.conditional_formatting.add("A3", FormulaRule(formula=['A3<>"ALL CHECKS PASS"'], font=Font(name=FONT, size=20, bold=True, color="C00000")))
put(c, "A4", f'=COUNTIF({RES},TRUE)&" of "&ROWS({RES})&" checks are TRUE."', font=F_SMALL)
note(c, f"A{CH_END + 1}", "The banner requires no FALSE and every result TRUE, so an error value in any check also fails it.")
presentation(c, "C00000", area=f"A1:E{CH_END + 1}")
c.freeze_panes = f"A{CH_H + 1}"

# ---------------------------------------------------------------- 9. README
rd = ws_readme
rd.sheet_properties.tabColor = "808080"
rd.column_dimensions["A"].width = 26
rd.column_dimensions["B"].width = 110
title(rd, "A1", "BUSM1042 Assignment 3: Q2 lead time and cancellation workbook")
lines = [
    ("Purpose", 'Business question 2: "How does booking lead time affect the likelihood of cancellation?" This '
                "workbook holds the raw data, the documented cleaning, formula-driven summary tables, self-checks "
                "and Figure 2 for the group report."),
    ("Date built", BUILD_DATE),
    ("Source dataset", "IBA-A3-Hotel_bookings_dataset-T32026.xlsx, sheet 'hotels' (119,390 bookings × 29 columns). "
                       "Data dictionary: IBA-A3-Dataset_description-T32026.xlsx (reference only)."),
    ("Calculation", "Formulas recalculate when the file opens (calculation set to automatic, full recalculation on "
                    "load). The first open can take a little while because Clean_Data holds about 1.2 million "
                    "row formulas."),
    ("Open and save", "Open this file in desktop Excel once, wait for the calculation to finish, then Save. This "
                      "stores the calculated values, so previews and Claude in Excel see numbers, not blanks."),
    ("Version", "v3, 23 October 2026: Figure 2 redesigned as a clustered column chart (all bookings and no-deposit "
                "bookings) with bookings and cancellations under each band; Table D and the figure checks rebuilt "
                "to match. All other sheets are unchanged from v1."),
    (None, None),
    ("SHEET MAP", None),
    ("README", "This page: purpose, sheet map, cleaning rules, band definitions and the group chart style guide."),
    ("Raw_Data", "Untouched copy of the source sheet (29 columns, same order, same rows). Protected without a password."),
    ("Clean_Data", "Raw values plus 11 derived columns (teal headers) in Excel Table tblClean; hover a teal header to "
                   "see its formula."),
    ("Table1_DataQuality", "Report Table 1 rebuilt with live counts, expected values and Match? flags, then the "
                           "copy-ready version for Word."),
    ("Q2_Summary", "Tables A–D: cancellation by lead-time band, 60-day comparison, robustness checks and chart data."),
    ("Figure2", "Native Excel clustered column chart for report Figure 2 (the only visual for Q2), with caption, "
                "note and source lines."),
    ("Checks", "TRUE/FALSE self-checks with an overall ALL CHECKS PASS banner."),
    ("Report_Figures", "Every number quoted in the report text, linked to its live cell and matched to the text."),
    (None, None),
    ("CLEANING RULES (Table 1)", None),
    ("No deletions", "All 119,390 bookings stay in every analysis. Problems are flagged, not removed."),
    ("Cancellation", "is_canceled holds Yes/No (not 0/1). cancel_flag = 1 when is_canceled = Yes; this includes "
                     "43,017 'Canceled' and 1,207 'No-Show' bookings (44,224 in total)."),
    ("Repeat guests", "repeat_flag = 1 when is_repeated_guest = Yes."),
    ("Blanks", "The source file stores missing values as the text NULL (country, agent, company) and NA (children). "
               "They are stored here as genuinely empty cells, in both Raw_Data and Clean_Data. This is the only "
               "change from the source, and it is what makes COUNTBLANK give 488 / 16,340 / 112,593 / 4."),
    ("Recodes", "country blank → \"Unknown\"; agent and company blank → \"None\" (no agent or company involved). "
                "agent and company are ID codes stored as whole numbers."),
    ("Dates", "96 arrival_date cells hold the text \"29 February 2023\" (an impossible date). They stay as text and "
              "arrival_date_valid = 0 flags them."),
    ("Other flags kept", "4 missing children, 180 bookings with zero guests, Average_Daily_Rate of −6.38 and 5,400: "
                         "kept unchanged because these fields enter no business question."),
    ("Duplicates (dup_flag)", "1 when a row is identical to an EARLIER row on all 28 columns except "
                              "booking_order_id, else 0 (32,239 rows). Computed in Python with pandas "
                              "df.duplicated(subset=all columns except booking_order_id, keep='first'), blanks "
                              "treated as equal, because a 28-column duplicate test is impractical as an Excel "
                              "formula over 119,390 rows. It is stored as a value; every other derived column is a "
                              "formula."),
    (None, None),
    ("LEAD-TIME BANDS", None),
    ("Bands (inclusive)", "0–7, 8–30, 31–90, 91–180, 181–365 and Over 365 (366 days or more)."),
    ("Suppression rule", "Rates based on fewer than 30 bookings are not plotted (group rule)."),
    ("Definitions", "No-deposit = deposit_type 'No Deposit'. Online TA = market_segment 'Online TA'."),
    (None, None),
    ("GROUP CHART STYLE GUIDE", None),
    ("Palette", "Okabe–Ito, colour-blind safe: #0072B2 blue (series 1), #009E73 green (series 2), #E69F00 orange "
                "(series 3), then #CC79A7 pink and #56B4E9 sky blue for extra series; grey #7F7F7F for context or "
                "benchmark lines."),
    ("Fonts", "Calibri 8–10 pt inside charts; Calibri 11 pt body, 12 pt bold table titles, 14 pt bold sheet titles."),
    ("Titles and captions", "No chart titles inside charts. The numbered caption goes above the figure in Word "
                            "(e.g. 'Figure 2: …'), and a source line goes under every figure."),
    ("Chart type", "Use columns for banded categories. Lines suggest a continuous trend, and their slopes would "
                   "mislead because the lead-time bands have unequal widths (8, 23, 60, 90 and 185 days, then "
                   "open-ended)."),
    ("Labels instead of axis", "When every bar carries its value label, hide the value axis and gridlines; keep the "
                               "scale fixed at 0–80% so bar heights stay honest."),
    ("Context in labels", "Put bookings and cancellations under each category so the chart shows volume as well as "
                          "rate."),
    ("Axes", "Rate axes start at 0%. Where a value axis is shown, light grey (#D9D9D9) major gridlines only; no 3-D, "
             "shadows or chart borders; white background."),
    ("Data labels", "One decimal place (0.0%), bold, in the series colour; no label may overlap another."),
    ("Legend", "Top, with plain-English series names."),
    ("Tables", "Navy (#1F3A5F) header with white bold text, thin #BFBFBF borders, #F2F2F2 banding, bold total rows, "
               "a numbered title above and a source/note line below."),
    ("Into Word", "Click the chart → Copy → in Word, Paste Special → Picture (Enhanced Metafile). Add the caption "
                  "above and the note and source lines below."),
    (None, None),
    ("FIGURE 2 DESIGN", None),
    ("Clustered columns", "Not lines: the bands have unequal widths. Bar length and position are the encodings "
                          "readers judge most accurately (Vidgen, Kirshner & Tan 2019, p. 194)."),
    ("Two series only", "A non-technical industry audience needs standard charts limited to the most important "
                        "variables (Vidgen, Kirshner & Tan 2019, pp. 198–199)."),
    ("No secondary axis", "A dual-axis combo was tested and rejected: it mixes counts and rates, and readers misread "
                          "the line height against the bars."),
    ("Value axis hidden", "Every bar carries a value label, so the axis adds ink but no information. The scale stays "
                          "fixed at 0–80%, so heights remain truthful."),
    ("Labels under the bars", "Volume and the deposit explanation are printed under the bars they describe, so no "
                              "reader has to look away from the chart to understand the over-365 drop."),
]
if VALUES_FALLBACK:
    lines.insert(4, ("Performance note", "To keep recalculation fast, repeat_flag, lead_over_60, country_clean, "
                                         "agent_clean, company_clean and guests_total are stored as values; each "
                                         "header's comment shows the formula used."))
r = 3
for k, v in lines:
    if k is None:
        r += 1
        continue
    if v is None:
        put(rd, f"A{r}", k, font=F_TTITLE)
        rd.row_dimensions[r].height = 18
    else:
        put(rd, f"A{r}", k, font=F_BOLD, align=Alignment(vertical="top"))
        put(rd, f"B{r}", v, align=A_WRAP)
        rd.row_dimensions[r].height = 15 * max(1, -(-len(v) // 120))
    r += 1
presentation(rd, "808080", area=f"A1:B{r}")

# ---------------------------------------------------------------- save
wb.active = 0
print(f"Saving {OUT} ...", flush=True)
wb.save(OUT)
print(f"Done in {time.time() - t0:.0f}s")
