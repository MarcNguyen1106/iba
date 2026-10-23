"""Verify the Q2 lead-time workbook (v3) three independent ways plus the golden expected values.

Usage:
  python3 -I verify.py recalc  <built.xlsx> <recalc_dir> <recalc.py>   # LibreOffice values-cached copy
  python3 -I verify.py check   <source.xlsx> <built.xlsx> <recalc.xlsx> <report.md>
  python3 -I verify.py render  <recalc.xlsx> <out_dir>                 # Figure2 / Q2_Summary previews

Method A = pandas on the ORIGINAL file; Method B = plain openpyxl row loop on the ORIGINAL file (no pandas);
Method C = the workbook's own Excel formulas, recalculated by LibreOffice and read with data_only=True.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import time
import zipfile
from collections import Counter
from datetime import datetime
from decimal import ROUND_HALF_UP, Decimal

N = 119390
BANDS = [("0–7", 0, 7), ("8–30", 8, 30), ("31–90", 31, 90), ("91–180", 91, 180),
         ("181–365", 181, 365), ("Over 365", 366, 100000)]
LABELS = [b[0] for b in BANDS]
OK = "✔ OK"

# ------------------------------------------------------------------ Section 7 golden reference
G_ALL = {"0–7": (19746, 1902, 9.63), "8–30": (18960, 5283, 27.86), "31–90": (29553, 11141, 37.70),
         "91–180": (26439, 11821, 44.71), "181–365": (21544, 11947, 55.45), "Over 365": (3148, 2130, 67.66),
         "Total": (119390, 44224, 37.04)}
G_ND = {"0–7": (19619, 1798, 9.16), "8–30": (18438, 4802, 26.04), "31–90": (27139, 8736, 32.19),
        "91–180": (22260, 7760, 34.86), "181–365": (15905, 6334, 39.82), "Over 365": (1280, 264, 20.62),
        "Total": (104641, 29694, 28.38)}
G_OTA = {"0–7": (8315, 767, 9.22), "8–30": (10923, 3494, 31.99), "31–90": (16454, 6520, 39.63),
         "91–180": (13233, 5889, 44.50), "181–365": (7462, 3984, 53.39), "Over 365": (16, 14, 87.50),
         "Total": (56403, 20668, None)}
G_NR = {"0–7": (121, 0.61), "8–30": (494, 2.61), "31–90": (2406, 8.14), "91–180": (4087, 15.46),
        "181–365": (5614, 26.06), "Over 365": (1865, 59.24), "Total": (14587, 12.22)}
G_SHB = {"0–7": 16.54, "8–30": 15.88, "31–90": 24.75, "91–180": 22.15, "181–365": 18.05, "Over 365": 2.64,
         "Total": 100.0}
G_SHC = {"0–7": 4.30, "8–30": 11.95, "31–90": 25.19, "91–180": 26.73, "181–365": 27.01, "Over 365": 4.82,
         "Total": 100.0}
G_B = {"nd_gt60": (50599, 17901, 35.38), "nd_le60": (54042, 11793, 21.82), "nd_rr": 1.62, "nd_more": 62,
       "all_gt60": (63714, None, 48.44), "all_le60": (55676, None, 23.99), "all_rr": 2.02}
G_C2 = (77536, 64.94, 34909, 78.94)
G_DD = {"0–7": (18302, 1541, 8.42), "8–30": (16317, 4122, 25.26), "31–90": (22668, 7204, 31.78),
        "91–180": (18171, 6310, 34.73), "181–365": (11134, 4378, 39.32), "Over 365": (559, 225, 40.25),
        "Total": (87151, 23780, 27.29)}
G_T1 = {"rows": 119390, "canc": 44224, "st_canc": 43017, "st_ns": 1207, "st_co": 75166, "mismatch": 0,
        "lt_blank": 0, "lt_min": 0, "lt_max": 737, "dups": 32239, "dup_go": 25313, "dup_go_sh_pct": 78.5,
        "go": 44030, "go_sh_pct": 36.9, "country": 488, "undef": 5, "agent": 16340, "company": 112593,
        "dates": 96, "children": 4, "zero": 180, "adr_min": -6.38, "adr_max": 5400, "nodep": 104641,
        "nonref": 14587, "refund": 162}


def r1(x):
    """Per cent, 1 dp, Decimal half-up (x is a fraction)."""
    return Decimal(repr(float(x) * 100)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def r2(x):
    return Decimal(repr(float(x) * 100)).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def g1(pct):
    return Decimal(str(pct)).quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)


def band_of(lt):
    if lt <= 7:
        return "0–7"
    elif lt <= 30:
        return "8–30"
    elif lt <= 90:
        return "31–90"
    elif lt <= 180:
        return "91–180"
    elif lt <= 365:
        return "181–365"
    else:
        return "Over 365"


# ================================================================== stage: recalc
def stage_recalc(built, recalc_dir, recalc_py):
    """Copy the build, add print areas to the two data sheets (so a PDF preview stays small), recalc in LibreOffice."""
    os.makedirs(recalc_dir, exist_ok=True)
    dst = os.path.abspath(os.path.join(recalc_dir, os.path.basename(built)))
    with zipfile.ZipFile(built) as zin, zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/workbook.xml":
                xml = data.decode("utf8")
                extra = ('<definedName name="_xlnm.Print_Area" localSheetId="1">Raw_Data!$A$1:$E$20</definedName>'
                         '<definedName name="_xlnm.Print_Area" localSheetId="2">Clean_Data!$A$1:$E$20</definedName>')
                if "<definedNames>" in xml:
                    xml = xml.replace("<definedNames>", "<definedNames>" + extra, 1)
                else:
                    xml = xml.replace("</sheets>", "</sheets><definedNames>" + extra + "</definedNames>", 1)
                data = xml.encode("utf8")
            zout.writestr(item, data)
    t = time.time()
    out = subprocess.run([sys.executable, os.path.abspath(recalc_py), dst, "3600"], capture_output=True, text=True,
                         cwd=os.path.dirname(recalc_py))
    secs = time.time() - t
    print(out.stdout[-3000:], out.stderr[-3000:])
    with open(os.path.join(recalc_dir, "recalc_time.txt"), "w") as fh:
        fh.write(f"{secs:.0f}\n{out.stdout}")
    print(f"LibreOffice recalculation + error scan: {secs:.0f}s")


# ================================================================== Method A: pandas
def method_a(src):
    import pandas as pd
    d = pd.read_excel(src, sheet_name="hotels", engine="openpyxl")
    assert d.shape == (N, 29)
    d["cf"] = (d["is_canceled"] == "Yes").astype(int)
    d["band"] = pd.cut(d["lead_time"], bins=[-1, 7, 30, 90, 180, 365, 10 ** 6], labels=LABELS).astype(str)
    d["dup"] = d.duplicated(subset=list(d.columns[1:29]), keep="first").astype(int)
    nd = d["deposit_type"] == "No Deposit"
    ota = nd & (d["market_segment"] == "Online TA")
    res = {}
    tot_b, tot_c = len(d), int(d["cf"].sum())
    for lab in LABELS + ["Total"]:
        m = (d["band"] == lab) if lab != "Total" else pd.Series(True, index=d.index)
        res[("A", lab, "n")] = int(m.sum())
        res[("A", lab, "c")] = int(d.loc[m, "cf"].sum())
        res[("A", lab, "rate")] = d.loc[m, "cf"].sum() / m.sum()
        res[("A", lab, "nd_n")] = int((m & nd).sum())
        res[("A", lab, "nd_c")] = int(d.loc[m & nd, "cf"].sum())
        res[("A", lab, "nd_rate")] = d.loc[m & nd, "cf"].sum() / (m & nd).sum()
        res[("A", lab, "ota_n")] = int((m & ota).sum())
        res[("A", lab, "ota_c")] = int(d.loc[m & ota, "cf"].sum())
        if lab != "Total":
            res[("A", lab, "ota_rate")] = d.loc[m & ota, "cf"].sum() / (m & ota).sum()
        nr = m & (d["deposit_type"] == "Non Refund")
        res[("A", lab, "nr_n")] = int(nr.sum())
        res[("A", lab, "nr_sh")] = nr.sum() / m.sum()
        res[("A", lab, "sh_b")] = m.sum() / tot_b
        res[("A", lab, "sh_c")] = d.loc[m, "cf"].sum() / tot_c
        dd = m & (d["dup"] == 0)
        res[("C3", lab, "n")] = int(dd.sum())
        res[("C3", lab, "c")] = int(d.loc[dd, "cf"].sum())
        res[("C3", lab, "rate")] = d.loc[dd, "cf"].sum() / dd.sum()
    for grp, base in (("nd", nd), ("all", pd.Series(True, index=d.index))):
        gt, le = base & (d["lead_time"] > 60), base & (d["lead_time"] <= 60)
        res[("B", grp, "gt_n")] = int(gt.sum())
        res[("B", grp, "gt_c")] = int(d.loc[gt, "cf"].sum())
        res[("B", grp, "le_n")] = int(le.sum())
        res[("B", grp, "le_c")] = int(d.loc[le, "cf"].sum())
        res[("B", grp, "gt_rate")] = d.loc[gt, "cf"].sum() / gt.sum()
        res[("B", grp, "le_rate")] = d.loc[le, "cf"].sum() / le.sum()
        res[("B", grp, "rr")] = res[("B", grp, "gt_rate")] / res[("B", grp, "le_rate")]
        res[("B", grp, "more")] = res[("B", grp, "rr")] - 1
    m = (d["lead_time"] >= 31) & (d["lead_time"] <= 365)
    res[("C2", "n")] = int(m.sum())
    res[("C2", "sh_b")] = m.sum() / tot_b
    res[("C2", "c")] = int(d.loc[m, "cf"].sum())
    res[("C2", "sh_c")] = d.loc[m, "cf"].sum() / tot_c
    # Table 1
    seg, chan = d["market_segment"], d["distribution_channel"]
    guests = d["adults"] + d["children"].fillna(0) + d["babies"]
    adr = d["Average_Daily_Rate"]
    go = seg.isin(["Groups", "Offline TA/TO"])
    t1 = {
        "rows": int(d["booking_order_id"].notna().sum()), "city": int((d["hotel_name"] == "City Hotel").sum()),
        "resort": int((d["hotel_name"] == "Resort Hotel").sum()),
        "yesno": int(d["is_canceled"].isin(["Yes", "No"]).sum()),
        "rep_yesno": int(d["is_repeated_guest"].isin(["Yes", "No"]).sum()),
        "canc": int((d["is_canceled"] == "Yes").sum()),
        "st_canc": int((d["reservation_status"] == "Canceled").sum()),
        "st_ns": int((d["reservation_status"] == "No-Show").sum()),
        "st_co": int((d["reservation_status"] == "Check-Out").sum()),
        "mismatch": int(((d["is_canceled"] == "Yes") & (d["reservation_status"] == "Check-Out")).sum()
                        + ((d["is_canceled"] == "No") & (d["reservation_status"] != "Check-Out")).sum()),
        "lt_blank": int(d["lead_time"].isna().sum()), "lt_neg": int((d["lead_time"] < 0).sum()),
        "lt_min": int(d["lead_time"].min()), "lt_max": int(d["lead_time"].max()),
        "lt_med": float(d["lead_time"].median()), "lt_mean": float(d["lead_time"].mean()),
        "dups": int(d["dup"].sum()), "dup_go": int((go & (d["dup"] == 1)).sum()),
        "go": int(go.sum()), "country": int(d["country"].isna().sum()),
        "seg_u": int((seg == "Undefined").sum()), "chan_u": int((chan == "Undefined").sum()),
        "undef": int(((seg == "Undefined") | (chan == "Undefined")).sum()),
        "agent": int(d["agent"].isna().sum()), "company": int(d["company"].isna().sum()),
        "ac_num": int((d["agent"].notna() | d["company"].notna()).sum()),
        "dates": int(d["arrival_date"].map(lambda v: not isinstance(v, datetime)).sum()),
        "children": int(d["children"].isna().sum()), "zero": int((guests == 0).sum()),
        "adr_min": float(adr.min()), "adr_max": float(adr.max()), "adr_neg": int((adr < 0).sum()),
        "adr_top": int((adr == adr.max()).sum()),
        "odd": int((d["children"].isna() | (guests == 0) | (adr < 0) | (adr == adr.max())).sum()),
        "nodep": int((d["deposit_type"] == "No Deposit").sum()),
        "nonref": int((d["deposit_type"] == "Non Refund").sum()),
        "refund": int((d["deposit_type"] == "Refundable").sum()),
    }
    t1["dup_go_sh"] = t1["dup_go"] / t1["dups"]
    t1["go_sh"] = t1["go"] / t1["rows"]
    for k, v in t1.items():
        res[("T1", k)] = v
    # per-row derived values for the Clean_Data formula check
    derived = {
        "cancel_flag": d["cf"].tolist(),
        "repeat_flag": (d["is_repeated_guest"] == "Yes").astype(int).tolist(),
        "lead_band": d["band"].tolist(),
        "lead_band_order": d["band"].map({l: i + 1 for i, l in enumerate(LABELS)}).astype(int).tolist(),
        "lead_over_60": (d["lead_time"] > 60).astype(int).tolist(),
        "country_clean": d["country"].fillna("Unknown").tolist(),
        "agent_clean": [("None" if pd.isna(v) else int(v)) for v in d["agent"]],
        "company_clean": [("None" if pd.isna(v) else int(v)) for v in d["company"]],
        "arrival_date_valid": d["arrival_date"].map(lambda v: 1 if isinstance(v, datetime) else 0).tolist(),
        "guests_total": guests.astype(int).tolist(),
        "dup_flag": d["dup"].tolist(),
    }
    return res, d, derived


# ================================================================== Method B: plain openpyxl loop, no pandas
def method_b(src):
    from openpyxl import load_workbook
    wb = load_workbook(src, read_only=True)
    ws = wb["hotels"]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    ix = {h: i for i, h in enumerate(hdr)}
    blank_tokens = {"country": "NULL", "agent": "NULL", "company": "NULL", "children": "NA"}
    cnt = Counter()
    seen = set()
    lt_vals = []
    adr_vals = []
    rows_raw = []
    for row in it:
        rows_raw.append(row)
        lt = row[ix["lead_time"]]
        lt_vals.append(lt)
        adr_vals.append(row[ix["Average_Daily_Rate"]])
    wb.close()
    adr_max = max(adr_vals)
    for row in rows_raw:
        lt = row[ix["lead_time"]]
        band = band_of(lt)
        canc = 1 if row[ix["is_canceled"]] == "Yes" else 0
        dep = row[ix["deposit_type"]]
        seg = row[ix["market_segment"]]
        chan = row[ix["distribution_channel"]]
        key = tuple(row[1:])
        is_dup = key in seen
        seen.add(key)

        def blank(f):
            return row[ix[f]] is None or row[ix[f]] == blank_tokens.get(f, object())
        for lab in (band, "Total"):
            cnt[("A", lab, "n")] += 1
            cnt[("A", lab, "c")] += canc
            if dep == "No Deposit":
                cnt[("A", lab, "nd_n")] += 1
                cnt[("A", lab, "nd_c")] += canc
                if seg == "Online TA":
                    cnt[("A", lab, "ota_n")] += 1
                    cnt[("A", lab, "ota_c")] += canc
            if dep == "Non Refund":
                cnt[("A", lab, "nr_n")] += 1
            if not is_dup:
                cnt[("C3", lab, "n")] += 1
                cnt[("C3", lab, "c")] += canc
        for grp, inc in (("nd", dep == "No Deposit"), ("all", True)):
            if inc:
                side = "gt" if lt > 60 else "le"
                cnt[("B", grp, side + "_n")] += 1
                cnt[("B", grp, side + "_c")] += canc
        if 31 <= lt <= 365:
            cnt[("C2", "n")] += 1
            cnt[("C2", "c")] += canc
        st = row[ix["reservation_status"]]
        children = 0 if blank("children") else row[ix["children"]]
        guests = row[ix["adults"]] + children + row[ix["babies"]]
        adr = row[ix["Average_Daily_Rate"]]
        t = cnt
        t[("T1", "rows")] += 1 if row[0] is not None else 0
        t[("T1", "city")] += row[ix["hotel_name"]] == "City Hotel"
        t[("T1", "resort")] += row[ix["hotel_name"]] == "Resort Hotel"
        t[("T1", "yesno")] += row[ix["is_canceled"]] in ("Yes", "No")
        t[("T1", "rep_yesno")] += row[ix["is_repeated_guest"]] in ("Yes", "No")
        t[("T1", "canc")] += canc
        t[("T1", "st_canc")] += st == "Canceled"
        t[("T1", "st_ns")] += st == "No-Show"
        t[("T1", "st_co")] += st == "Check-Out"
        t[("T1", "mismatch")] += (canc == 1 and st == "Check-Out") or (canc == 0 and st != "Check-Out")
        t[("T1", "lt_blank")] += lt is None
        t[("T1", "lt_neg")] += lt < 0
        t[("T1", "dups")] += is_dup
        t[("T1", "dup_go")] += is_dup and seg in ("Groups", "Offline TA/TO")
        t[("T1", "go")] += seg in ("Groups", "Offline TA/TO")
        t[("T1", "country")] += blank("country")
        t[("T1", "seg_u")] += seg == "Undefined"
        t[("T1", "chan_u")] += chan == "Undefined"
        t[("T1", "undef")] += seg == "Undefined" or chan == "Undefined"
        t[("T1", "agent")] += blank("agent")
        t[("T1", "company")] += blank("company")
        t[("T1", "ac_num")] += (not blank("agent")) or (not blank("company"))
        t[("T1", "dates")] += not isinstance(row[ix["arrival_date"]], datetime)
        t[("T1", "children")] += blank("children")
        t[("T1", "zero")] += guests == 0
        t[("T1", "adr_neg")] += adr < 0
        t[("T1", "adr_top")] += adr == adr_max
        t[("T1", "odd")] += blank("children") or guests == 0 or adr < 0 or adr == adr_max
        t[("T1", "nodep")] += dep == "No Deposit"
        t[("T1", "nonref")] += dep == "Non Refund"
        t[("T1", "refund")] += dep == "Refundable"
    res = {k: int(v) for k, v in cnt.items()}
    srt = sorted(lt_vals)
    res[("T1", "lt_min")], res[("T1", "lt_max")] = srt[0], srt[-1]
    res[("T1", "lt_med")] = float((srt[N // 2 - 1] + srt[N // 2]) / 2)
    res[("T1", "lt_mean")] = sum(lt_vals) / N
    res[("T1", "adr_min")], res[("T1", "adr_max")] = float(min(adr_vals)), float(adr_max)
    res[("T1", "dup_go_sh")] = res[("T1", "dup_go")] / res[("T1", "dups")]
    res[("T1", "go_sh")] = res[("T1", "go")] / res[("T1", "rows")]
    tb, tc = res[("A", "Total", "n")], res[("A", "Total", "c")]
    for lab in LABELS + ["Total"]:
        res[("A", lab, "rate")] = res[("A", lab, "c")] / res[("A", lab, "n")]
        res[("A", lab, "nd_rate")] = res[("A", lab, "nd_c")] / res[("A", lab, "nd_n")]
        if lab != "Total":
            res[("A", lab, "ota_rate")] = res[("A", lab, "ota_c")] / res[("A", lab, "ota_n")]
        res[("A", lab, "nr_sh")] = res[("A", lab, "nr_n")] / res[("A", lab, "n")]
        res[("A", lab, "sh_b")] = res[("A", lab, "n")] / tb
        res[("A", lab, "sh_c")] = res[("A", lab, "c")] / tc
        res[("C3", lab, "rate")] = res[("C3", lab, "c")] / res[("C3", lab, "n")]
    for grp in ("nd", "all"):
        res[("B", grp, "gt_rate")] = res[("B", grp, "gt_c")] / res[("B", grp, "gt_n")]
        res[("B", grp, "le_rate")] = res[("B", grp, "le_c")] / res[("B", grp, "le_n")]
        res[("B", grp, "rr")] = res[("B", grp, "gt_rate")] / res[("B", grp, "le_rate")]
        res[("B", grp, "more")] = res[("B", grp, "rr")] - 1
    res[("C2", "sh_b")] = res[("C2", "n")] / tb
    res[("C2", "sh_c")] = res[("C2", "c")] / tc
    return res, rows_raw, hdr


# ================================================================== Method C: workbook formulas (LibreOffice values)
def find_row(ws, text, col=1, start=1, stop=200):
    for r in range(start, stop):
        v = ws.cell(row=r, column=col).value
        if isinstance(v, str) and v.startswith(text):
            return r
    raise KeyError(text)


def method_c(recalc):
    from openpyxl import load_workbook
    wb = load_workbook(recalc, data_only=True)
    q = wb["Q2_Summary"]
    res, cells = {}, {}
    ta0 = find_row(q, "Table A:") + 2
    for i, lab in enumerate(LABELS + ["Total"]):
        r = ta0 + i
        assert (q[f"A{r}"].value == lab), (q[f"A{r}"].value, lab)
        for key, col in (("n", "D"), ("c", "E"), ("rate", "F"), ("nd_n", "G"), ("nd_c", "H"), ("nd_rate", "I"),
                         ("ota_n", "J"), ("ota_c", "K"), ("ota_rate", "L"), ("nr_n", "M"), ("nr_sh", "N"),
                         ("sh_b", "O"), ("sh_c", "P")):
            if lab == "Total" and key == "ota_rate":
                continue
            res[("A", lab, key)] = q[f"{col}{r}"].value
            cells[("A", lab, key)] = f"Q2_Summary!{col}{r}"
    b1 = find_row(q, "Table B:") + 2
    for grp, r in (("nd", b1), ("all", b1 + 2)):
        res[("B", grp, "gt_n")], res[("B", grp, "gt_c")], res[("B", grp, "gt_rate")] = \
            q[f"C{r}"].value, q[f"D{r}"].value, q[f"E{r}"].value
        res[("B", grp, "le_n")], res[("B", grp, "le_c")], res[("B", grp, "le_rate")] = \
            q[f"C{r + 1}"].value, q[f"D{r + 1}"].value, q[f"E{r + 1}"].value
        res[("B", grp, "rr")], res[("B", grp, "more")] = q[f"F{r}"].value, q[f"G{r}"].value
    c10 = find_row(q, "C1.") + 2
    for i, lab in enumerate(LABELS + ["Total"]):
        res[("C1", lab, "countif")] = q[f"C{c10 + i}"].value
        res[("C1", lab, "diff")] = q[f"D{c10 + i}"].value
    c2 = find_row(q, "C2.") + 2
    res[("C2", "n")], res[("C2", "sh_b")], res[("C2", "c")], res[("C2", "sh_c")] = \
        q[f"B{c2}"].value, q[f"C{c2}"].value, q[f"D{c2}"].value, q[f"E{c2}"].value
    c30 = find_row(q, "C3.") + 2
    for i, lab in enumerate(LABELS + ["Total"]):
        r = c30 + i
        res[("C3", lab, "n")], res[("C3", lab, "c")], res[("C3", lab, "rate")] = \
            q[f"B{r}"].value, q[f"C{r}"].value, q[f"D{r}"].value
    d0 = find_row(q, "Table D:") + 2
    res["D"] = [[q.cell(row=d0 + i, column=j).value for j in range(1, 5)] for i in range(6)]
    res["D_hdr"] = [q.cell(row=d0 - 1, column=j).value for j in range(1, 5)]
    res["D_rows"] = (d0 - 2, d0 - 1, d0, d0 + 5, d0 + 6)  # title, header, first, last, note
    res["D_note"] = q.cell(row=d0 + 6, column=1).value
    t = wb["Table1_DataQuality"]
    t1map = {}
    for r in range(5, 60):
        lab = t[f"B{r}"].value
        if lab is None or t[f"A{r}"].value is None or not isinstance(t[f"A{r}"].value, int):
            continue
        t1map[lab] = (t[f"D{r}"].value, t[f"E{r}"].value, t[f"F{r}"].value, r)
    res["T1rows"] = t1map
    pres_top = find_row(t, "Table 1: Data quality checks and treatment", start=5, stop=120) + 2
    res["T1pres"] = [[t.cell(row=r, column=j).value for j in range(1, 6)] for r in range(pres_top, pres_top + 10)]
    res["T1src"] = t.cell(row=pres_top + 10, column=1).value
    ch = wb["Checks"]
    res["banner"] = ch["A3"].value
    res["banner_count"] = ch["A4"].value
    res["checks"] = []
    for r in range(6, 80):
        if ch[f"B{r}"].value is None or not isinstance(ch[f"A{r}"].value, int):
            continue
        res["checks"].append((ch[f"A{r}"].value, ch[f"B{r}"].value, ch[f"C{r}"].value, ch[f"D{r}"].value,
                              ch[f"E{r}"].value))
    rf = wb["Report_Figures"]
    res["rf"] = []
    for r in range(5, 40):
        if rf[f"A{r}"].value is None or rf[f"E{r}"].value is None:
            continue
        res["rf"].append((rf[f"A{r}"].value, rf[f"C{r}"].value, rf[f"D{r}"].value, rf[f"E{r}"].value))
    f = wb["Figure2"]
    res["fig_caption"] = f["A1"].value
    res["fig_note"], res["fig_source"] = f["A24"].value, f["A25"].value
    res["fig_print_area"] = f.print_area
    wb.close()
    return res, cells


# ================================================================== structural checks
def same(a, b):
    if a is None and b is None:
        return True
    if isinstance(a, bool) or isinstance(b, bool):
        return a == b
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return float(a) == float(b)
    return a == b


def stage_check(src, built, recalc, report):
    lines, results = [], []
    t_start = time.time()

    def rec(section, name, a, b, c, exp, ok):
        results.append((section, name, a, b, c, exp, "PASS" if ok else "FAIL"))

    print("Method A (pandas) ...", flush=True)
    A, df, derived = method_a(src)
    print("Method B (openpyxl loop) ...", flush=True)
    B, raw_rows, hdr = method_b(src)
    print("Method C (recalculated workbook) ...", flush=True)
    Cv, cells = method_c(recalc)

    def fmt(v):
        if isinstance(v, float):
            return f"{v:.6f}".rstrip("0").rstrip(".")
        return str(v)

    # ---- Table A / B / C : A vs B vs C, exact counts, rates within 1e-12
    for key in sorted(k for k in A if isinstance(k, tuple) and k[0] in ("A", "B", "C2", "C3")):
        a, b, c = A[key], B.get(key), Cv.get(key)
        if isinstance(a, int):
            ok = a == b == c
        else:
            ok = b is not None and c is not None and abs(a - b) < 1e-12 and abs(a - c) < 1e-12
        rec("A/B/C agreement", " | ".join(map(str, key)), fmt(a), fmt(b), fmt(c), "", ok)
    # C1: COUNTIF(lead_band) equals band bookings
    for lab in LABELS + ["Total"]:
        c = Cv[("C1", lab, "countif")]
        rec("Table C1 (lead_band column)", f"{lab}: COUNTIF(lead_band)", A[("A", lab, "n")], B[("A", lab, "n")], c,
            "diff 0", c == A[("A", lab, "n")] and Cv[("C1", lab, "diff")] == 0)

    # ---- expected values (Section 7)
    def exp_count(name, key, val):
        ok = A[key] == B[key] == Cv[key] == val
        rec("Expected (Section 7)", name, A[key], B[key], Cv[key], val, ok)

    def exp_rate(name, key, pct):
        # The three methods must agree at 1 dp (Decimal half-up), and the value must equal the golden figure to
        # the precision the golden figure is given (2 dp, so within ±0.005 percentage points). Re-rounding a 2-dp
        # golden value to 1 dp would double-round (e.g. 11.95 -> 12.0 although the true value 11.946% is 11.9%).
        ok = r1(A[key]) == r1(B[key]) == r1(Cv[key]) and abs(float(Cv[key]) * 100 - pct) <= 0.005 + 1e-9
        rec("Expected (Section 7)", name, f"{r1(A[key])}% ({float(A[key]) * 100:.4f})",
            f"{r1(B[key])}% ({float(B[key]) * 100:.4f})", f"{r1(Cv[key])}% ({float(Cv[key]) * 100:.4f})",
            f"{pct}%", ok)

    for lab in LABELS + ["Total"]:
        n, c, p = G_ALL[lab]
        exp_count(f"{lab} all bookings", ("A", lab, "n"), n)
        exp_count(f"{lab} all cancelled", ("A", lab, "c"), c)
        exp_rate(f"{lab} all rate", ("A", lab, "rate"), p)
        n, c, p = G_ND[lab]
        exp_count(f"{lab} no-deposit bookings", ("A", lab, "nd_n"), n)
        exp_count(f"{lab} no-deposit cancelled", ("A", lab, "nd_c"), c)
        exp_rate(f"{lab} no-deposit rate", ("A", lab, "nd_rate"), p)
        n, c, p = G_OTA[lab]
        exp_count(f"{lab} no-deposit Online TA bookings", ("A", lab, "ota_n"), n)
        exp_count(f"{lab} no-deposit Online TA cancelled", ("A", lab, "ota_c"), c)
        if p is not None:
            exp_rate(f"{lab} no-deposit Online TA rate", ("A", lab, "ota_rate"), p)
        n, p = G_NR[lab]
        exp_count(f"{lab} Non Refund bookings", ("A", lab, "nr_n"), n)
        exp_rate(f"{lab} Non Refund share", ("A", lab, "nr_sh"), p)
        exp_rate(f"{lab} share of bookings", ("A", lab, "sh_b"), G_SHB[lab])
        exp_rate(f"{lab} share of cancellations", ("A", lab, "sh_c"), G_SHC[lab])
        n, c, p = G_DD[lab]
        exp_count(f"{lab} de-duplicated bookings", ("C3", lab, "n"), n)
        exp_count(f"{lab} de-duplicated cancelled", ("C3", lab, "c"), c)
        exp_rate(f"{lab} de-duplicated rate", ("C3", lab, "rate"), p)
    for grp, gt, le in (("nd", "nd_gt60", "nd_le60"), ("all", "all_gt60", "all_le60")):
        exp_count(f"{grp} >60 bookings", ("B", grp, "gt_n"), G_B[gt][0])
        exp_count(f"{grp} ≤60 bookings", ("B", grp, "le_n"), G_B[le][0])
        if G_B[gt][1] is not None:
            exp_count(f"{grp} >60 cancelled", ("B", grp, "gt_c"), G_B[gt][1])
            exp_count(f"{grp} ≤60 cancelled", ("B", grp, "le_c"), G_B[le][1])
        exp_rate(f"{grp} >60 rate", ("B", grp, "gt_rate"), G_B[gt][2])
        exp_rate(f"{grp} ≤60 rate", ("B", grp, "le_rate"), G_B[le][2])
        key = ("B", grp, "rr")
        want = Decimal(str(G_B[grp + "_rr"]))
        vals = [Decimal(repr(float(X[key]))).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP) for X in (A, B, Cv)]
        rec("Expected (Section 7)", f"{grp} relative risk", *vals, want, all(v == want for v in vals))
    key = ("B", "nd", "more")
    vals = [Decimal(repr(float(X[key]) * 100)).quantize(Decimal("1"), rounding=ROUND_HALF_UP) for X in (A, B, Cv)]
    rec("Expected (Section 7)", "nd % more likely", *[f"{v}%" for v in vals], "62%", all(v == 62 for v in vals))
    exp_count("31–365 bookings", ("C2", "n"), G_C2[0])
    exp_rate("31–365 share of bookings", ("C2", "sh_b"), G_C2[1])
    exp_count("31–365 cancelled", ("C2", "c"), G_C2[2])
    exp_rate("31–365 share of cancellations", ("C2", "sh_c"), G_C2[3])

    # ---- Table 1 counts: A vs B vs C vs Section 3.4 expected
    t1labels = {
        "rows": "Rows in file (bookings)", "city": "City Hotel bookings", "resort": "Resort Hotel bookings",
        "yesno": "is_canceled coded as text Yes/No", "rep_yesno": "is_repeated_guest coded as text Yes/No",
        "canc": "is_canceled = Yes", "st_canc": "…of which reservation_status = Canceled",
        "st_ns": "…of which reservation_status = No-Show", "st_co": "reservation_status = Check-Out",
        "mismatch": "is_canceled vs reservation_status mismatches", "lt_blank": "lead_time blanks",
        "lt_neg": "lead_time negative values", "lt_min": "lead_time minimum (days)",
        "lt_max": "lead_time maximum (days)", "lt_med": "lead_time median (days)", "lt_mean": "lead_time mean (days)",
        "dups": "Duplicate rows (dup_flag = 1)", "dup_go": "Duplicates that are Groups or Offline TA/TO",
        "dup_go_sh": "…as a share of all duplicates", "go": "Groups + Offline TA/TO bookings",
        "go_sh": "…as a share of all bookings", "country": "Missing country",
        "seg_u": "market_segment = Undefined", "chan_u": "distribution_channel = Undefined",
        "undef": "Rows with segment or channel = Undefined", "agent": "Blank agent", "company": "Blank company",
        "ac_num": "agent or company stored as numbers (non-blank)",
        "dates": "Invalid arrival dates (arrival_date_valid = 0)", "children": "Missing children",
        "zero": "Zero guests (guests_total = 0)", "adr_min": "Average_Daily_Rate minimum",
        "adr_max": "Average_Daily_Rate maximum", "adr_neg": "Bookings with negative Average_Daily_Rate",
        "adr_top": "Bookings at the maximum Average_Daily_Rate",
        "odd": "Distinct bookings with missing children, zero guests or an extreme rate",
        "nodep": "deposit_type = No Deposit", "nonref": "deposit_type = Non Refund",
        "refund": "deposit_type = Refundable",
    }
    t1rows = Cv["T1rows"]
    assert len(t1rows) == len(t1labels), (len(t1rows), len(t1labels))
    for k, lab in t1labels.items():
        a, b = A[("T1", k)], B[("T1", k)]
        c, typed, match, row = t1rows[lab]
        ok = abs(a - b) < 1e-9 and abs(a - c) < 1e-9 and abs(round(a, 4) - round(typed, 4)) < 1e-12 and match == OK
        exp = typed
        if k in G_T1:
            ok = ok and abs(a - G_T1[k]) < 1e-9
            exp = G_T1[k]
        if k + "_pct" in G_T1:
            ok = ok and r1(a) == g1(G_T1[k + "_pct"])
            exp = f"{G_T1[k + '_pct']}%"
        rec("Table 1 counts", f"{lab} (row {row})", fmt(a), fmt(b), fmt(c), exp, ok)

    # ---- Table D (feeds Figure 2): displayed labels character for character, values, check column
    D = Cv["D"]
    rec("Table D (chart data)", "position: title / header / data / note rows", "", "", str(Cv["D_rows"]),
        "(49, 50, 51, 56, 57)", Cv["D_rows"] == (49, 50, 51, 56, 57))
    nr_over = int(Decimal(str(G_NR["Over 365"][1])).quantize(Decimal("1"), rounding=ROUND_HALF_UP))  # 59.24 -> 59
    for i, lab in enumerate(LABELS):
        row = D[i]
        n, c, _ = G_ALL[lab]  # expected label built from the golden values, not from the workbook
        want_cat = f"{lab}\n{n:,} bookings\n{c:,} cancelled" + (f"\n{nr_over}% non-refundable"
                                                                  if lab == "Over 365" else "")
        rec("Table D (chart data)", f"{lab}: displayed axis label (exact characters)", "", "", repr(row[0]),
            repr(want_cat), row[0] == want_cat)
        ok = (abs(row[1] - A[("A", lab, "rate")]) < 1e-12 and abs(row[1] - B[("A", lab, "rate")]) < 1e-12
              and abs(row[2] - A[("A", lab, "nd_rate")]) < 1e-12 and abs(row[2] - B[("A", lab, "nd_rate")]) < 1e-12
              and r1(row[1]) == g1(G_ALL[lab][2]) and abs(float(row[1]) * 100 - G_ALL[lab][2]) <= 0.005 + 1e-9
              and r1(row[2]) == g1(G_ND[lab][2]) and abs(float(row[2]) * 100 - G_ND[lab][2]) <= 0.005 + 1e-9)
        rec("Table D (chart data)", f"{lab}: all / no-deposit rate", f"{r1(A[('A', lab, 'rate')])}% / "
            f"{r1(A[('A', lab, 'nd_rate')])}%", f"{r1(B[('A', lab, 'rate')])}% / {r1(B[('A', lab, 'nd_rate')])}%",
            f"{r1(row[1])}% / {r1(row[2])}%", f"{G_ALL[lab][2]}% / {G_ND[lab][2]}%", ok)
        rec("Table D (chart data)", f"{lab}: 'Check: equals Table A?' cell", "", "", row[3], True, row[3] is True)
    want_hdr = ["Category label (axis)", "Cancellation rate, all bookings", "Cancellation rate, no-deposit bookings",
                "Check: equals Table A?"]
    rec("Table D (chart data)", "headers (B50 and C50 are the legend entries)", "", "", " | ".join(Cv["D_hdr"]),
        " | ".join(want_hdr), Cv["D_hdr"] == want_hdr)
    want_dnote = ("Each bar is the share of bookings in that band that were cancelled (is_canceled = Yes, including "
                  "no-shows). Labels show each band's bookings and cancellations. Source: Table A above.")
    rec("Table D (chart data)", "A57 note", "", "", Cv["D_note"], want_dnote, Cv["D_note"] == want_dnote)
    for num, lab, left, right, res_ in Cv["checks"]:
        rec("Checks sheet", f"#{num} {lab}", "", "", f"{fmt(left)} vs {fmt(right)} → {res_}", "TRUE", res_ is True)
    rec("Checks sheet", "Overall banner", "", "", Cv["banner"], "ALL CHECKS PASS", Cv["banner"] == "ALL CHECKS PASS")
    for lab, live, exp, match in Cv["rf"]:
        rec("Report_Figures", lab, "", "", f"{fmt(live)} → {match}", fmt(exp), match == OK)
    want_cap = "Figure 2: Cancellation rate by booking lead time, all bookings and no-deposit bookings"
    rec("Figure 2 text", "Caption (A1)", "", "", Cv["fig_caption"], want_cap, Cv["fig_caption"] == want_cap)
    want_note = ("Note: Each bar is the share of bookings in that lead-time band that were cancelled, including "
                 "no-shows. Labels under each band give its bookings and cancellations. Beyond 365 days, 59.2% of "
                 "bookings are non-refundable, so most of that band sits outside the no-deposit bar.")
    want_src = "Source: Group analysis of the supplied hotel bookings dataset (n = 119,390)."
    rec("Figure 2 text", "Note line (live formula)", "", "", Cv["fig_note"], want_note, Cv["fig_note"] == want_note)
    rec("Figure 2 text", "Source line (live formula)", "", "", Cv["fig_source"], want_src, Cv["fig_source"] == want_src)
    pa = str(Cv["fig_print_area"])
    rec("Figure 2 text", "Print area", "", "", pa, "A1:J27", pa.replace("$", "").endswith("A1:J27"))

    # Table 1 presentation text vs the report text (docx), hyphen/quote-normalised
    report_t1 = [
        ("is_canceled and is_repeated_guest are coded Yes/No, not 0/1 as the data dictionary states", "119,390"),
        ("is_canceled = Yes covers 43,017 cancellations and 1,207 no-shows, and it agrees with reservation_status "
         "for every booking", "44,224"),
        ("lead_time has no blanks or negative values (0 to 737 days) but is right-skewed (median 69, mean 104 days)",
         "0"),
        ("Rows identical on all 28 variables but with different booking IDs", "32,239"),
        ("Missing country", "488"),
        ('market_segment (2) or distribution_channel (5) recorded as "Undefined"', "5"),
        ("Blank agent and company", "16,340 / 112,593"),
        ("agent and company are stored as numbers but are ID codes", "119,390"),
        ('arrival_date holds full dates, although the dictionary says "Year of arrival". The original 2015–2017 '
         'dates were moved forward seven years, turning 29 February 2016 into the impossible text date '
         '"29 February 2023"', "96"),
        ("Missing children (4), no guests recorded (180), Average_Daily_Rate of −6.38 (1) and 5,400 (1)", "186"),
    ]
    dedup_text = "Kept: 78.5% are Groups or Offline TA/TO bookings (36.9% of all bookings), so they most likely " \
                 "record separate rooms in multi-room bookings. The Q2 pattern holds without them (8.4% at 0–7 " \
                 "days rising to 39.3% at 181–365 days)"

    def norm(s):
        return str(s).replace("−", "-").replace("“", '"').replace("”", '"')
    for (want_txt, want_n), got in zip(report_t1, Cv["T1pres"]):
        got_n = got[2] if isinstance(got[2], str) else f"{got[2]:,}"
        ok = norm(got[1]) == norm(want_txt) and got_n == want_n
        rec("Table 1 text vs report", want_txt[:60] + "…", "", "", f"{got[1][:60]}… | {got_n}", want_n, ok)
    rec("Table 1 text vs report", "Duplicate-row treatment text (live %)", "", "", Cv["T1pres"][3][4][:70] + "…",
        "matches report", norm(Cv["T1pres"][3][4]) == norm(dedup_text))

    # ---- structural checks on the built (formula) workbook
    print("Structural checks ...", flush=True)
    from openpyxl import load_workbook
    wbf = load_workbook(built, read_only=True)
    rawws = wbf["Raw_Data"]
    tokens = {"country": "NULL", "agent": "NULL", "company": "NULL", "children": "NA"}
    tok_idx = {hdr.index(k): v for k, v in tokens.items()}
    pd_vals = df[hdr].astype(object).where(df[hdr].notna(), None).values.tolist()
    mism_src, mism_pd, null_mapped, nan_strings = 0, 0, 0, 0
    raw_built = []
    for i, row in enumerate(rawws.iter_rows(values_only=True)):
        row = list(row)[:29]
        if i == 0:
            hdr_ok = row == hdr
            continue
        raw_built.append(row)
        src = raw_rows[i - 1]
        for j in range(29):
            v, s = row[j], src[j]
            if isinstance(v, str) and v.lower() in ("nan", "null"):
                nan_strings += 1
            if not same(v, s):
                if j in tok_idx and s == tok_idx[j] and v is None:
                    null_mapped += 1
                else:
                    mism_src += 1
            if not same(v, pd_vals[i - 1][j]):
                mism_pd += 1
    rec("Structure", "Raw_Data header = source header", "", "", hdr_ok, True, hdr_ok)
    rec("Structure", "Raw_Data rows", "", "", len(raw_built), N, len(raw_built) == N)
    rec("Structure", "Raw_Data vs pandas-loaded source, 119,390 × 29 cells (NaN = blank, int = float)", "", "",
        f"{mism_pd} mismatches", 0, mism_pd == 0)
    rec("Structure", "Raw_Data vs source cell values; differences other than NULL/NA text → blank", "", "",
        f"{mism_src} mismatches ({null_mapped:,} NULL/NA cells now blank)", "0 (129,425 blanks)",
        mism_src == 0 and null_mapped == 488 + 16340 + 112593 + 4)
    clws = wbf["Clean_Data"]
    cl_mism, form_bad, dupsum = 0, 0, 0
    cl_hdr = None
    for i, row in enumerate(clws.iter_rows(values_only=True)):
        if i == 0:
            cl_hdr = list(row)
            continue
        if any(not same(a, b) for a, b in zip(row[:29], raw_built[i - 1])):
            cl_mism += 1
        f = row[29:39]
        if not all(isinstance(x, str) and x.startswith("=") and f"{i + 1}" in x for x in f):
            form_bad += 1
        dupsum += row[39]
    rec("Structure", "Clean_Data first 29 columns = Raw_Data", "", "", f"{cl_mism} rows differ", 0, cl_mism == 0)
    rec("Structure", "Clean_Data derived columns are row formulas (10 per row)", "", "", f"{form_bad} rows not formulas",
        0, form_bad == 0)
    rec("Structure", "Clean_Data dup_flag values sum", "", "", dupsum, 32239, dupsum == 32239)
    want_hdr = hdr + ["cancel_flag", "repeat_flag", "lead_band", "lead_band_order", "lead_over_60", "country_clean",
                      "agent_clean", "company_clean", "arrival_date_valid", "guests_total", "dup_flag"]
    rec("Structure", "Clean_Data headers", "", "", len(cl_hdr), 40, cl_hdr == want_hdr)
    # "nan" strings anywhere in the formula workbook (all sheets)
    for ws in wbf.worksheets:
        if ws.title in ("Raw_Data",):
            continue
        for row in ws.iter_rows(values_only=True):
            for v in row:
                if isinstance(v, str) and v.strip().lower() in ("nan", "null", "na"):
                    nan_strings += 1
    rec("Structure", "Cells holding 'nan' / 'NaN' / 'NULL' / 'NA' text (all sheets)", "", "", nan_strings, 0,
        nan_strings == 0)
    names = list(wbf.sheetnames)
    want_names = ["README", "Raw_Data", "Clean_Data", "Table1_DataQuality", "Q2_Summary", "Figure2", "Checks",
                  "Report_Figures"]
    rec("Structure", "Sheet order", "", "", ", ".join(names), ", ".join(want_names), names == want_names)
    wbf.close()

    # Clean_Data formula results (LibreOffice) vs Python, every row
    wbr = load_workbook(recalc, read_only=True, data_only=True)
    der_names = list(derived.keys())
    bad = Counter()
    for i, row in enumerate(wbr["Clean_Data"].iter_rows(min_row=2, values_only=True)):
        vals = row[29:40]
        for name, v in zip(der_names, vals):
            want = derived[name][i]
            if not same(v, want):
                bad[name] += 1
    wbr.close()
    for name in der_names:
        rec("Clean_Data formula results vs Python (all 119,390 rows)", name, "", "", f"{bad[name]} mismatches", 0,
            bad[name] == 0)

    # formula-text audit: Excel-2019 safe functions, no structured refs, no typed numbers in summary tables
    wbf = load_workbook(built)
    banned = re.compile(r"\b(IFS|LET|LAMBDA|XLOOKUP|XMATCH|FILTER|UNIQUE|SORT|SORTBY|SEQUENCE|MAXIFS|MINIFS|"
                        r"TEXTJOIN|CONCAT|SWITCH)\s*\(")
    n_form, n_banned, n_struct = 0, 0, 0
    for ws in wbf.worksheets:
        if ws.title in ("Raw_Data", "Clean_Data"):
            continue
        for row in ws.iter_rows():
            for cell in row:
                if isinstance(cell.value, str) and cell.value.startswith("="):
                    n_form += 1
                    n_banned += bool(banned.search(cell.value))
                    n_struct += "[@" in cell.value or "tblClean[" in cell.value
    rec("Structure", f"Banned (post-2019 / dynamic-array) functions in {n_form} summary formulas", "", "", n_banned, 0,
        n_banned == 0)
    rec("Structure", "Structured [@col] references", "", "", n_struct, 0, n_struct == 0)
    q = wbf["Q2_Summary"]
    typed = []
    for row in q.iter_rows(min_row=6, max_row=60, min_col=4, max_col=16):
        for cell in row:
            if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                typed.append(cell.coordinate)
    rec("Structure", "Typed numbers inside Q2_Summary result columns (D:P)", "", "", ", ".join(typed) or "none", "none",
        not typed)
    calc = wbf.calculation
    rec("Structure", "fullCalcOnLoad set", "", "", calc.fullCalcOnLoad, True, bool(calc.fullCalcOnLoad))
    rec("Structure", "Named ranges", "", "", ", ".join(sorted(wbf.defined_names.keys())), "rng_* defined",
        all(n in wbf.defined_names for n in ("rng_lead", "rng_cancel", "rng_deposit", "rng_segment", "rng_band",
                                              "rng_dup")))
    tbl_ok = "tblClean" in wbf["Clean_Data"].tables
    rec("Structure", "Clean_Data is Excel Table tblClean", "", "", wbf["Clean_Data"].tables.get("tblClean").ref
        if tbl_ok else None, "A1:AN119391", tbl_ok and wbf["Clean_Data"].tables["tblClean"].ref == "A1:AN119391")
    rec("Structure", "Raw_Data protected (no password)", "", "", wbf["Raw_Data"].protection.sheet, True,
        wbf["Raw_Data"].protection.sheet and not wbf["Raw_Data"].protection.password)
    wbf.close()

    # chart XML (tag matching ignores namespace prefixes: openpyxl writes no "c:" prefix)
    z = zipfile.ZipFile(built)
    charts = sorted(n for n in z.namelist() if re.match(r"xl/charts/chart\d+\.xml$", n))
    rec("Chart XML (Figure 2)", "exactly one chart in the workbook", "", "", ", ".join(charts), 1, len(charts) == 1)
    xml = re.sub(r"<(/?)[a-zA-Z0-9]+:", r"<\1", z.read(charts[0]).decode("utf8"))  # drop every prefix
    xml = re.sub(r' xmlns(:\w+)?="[^"]*"', "", xml)

    def block(tag):
        m = re.search(rf"<{tag}>(.*?)</{tag}>", xml, flags=re.S)
        return m.group(1) if m else ""
    val_ax, cat_ax, legend = block("valAx"), block("catAx"), block("legend")
    sers = re.findall(r"<ser>(.*?)</ser>", xml, flags=re.S)
    chart_lines = []
    for sr in sers:
        tx = re.search(r"<tx><strRef><f>(.*?)</f>", sr)
        cat = re.search(r"<cat><strRef><f>(.*?)</f>", sr)
        val = re.search(r"<val><numRef><f>(.*?)</f>", sr)
        fill = re.search(r'<spPr><solidFill><srgbClr val="(\w+)"/></solidFill><ln><noFill/>', sr)
        dl = re.search(r"<dLbls>(.*?)</dLbls>", sr, flags=re.S)
        dl = dl.group(1) if dl else ""
        chart_lines.append(dict(
            tx=tx.group(1) if tx else None, cat=cat.group(1) if cat else None, val=val.group(1) if val else None,
            fill=fill.group(1) if fill else None,
            pos=re.findall(r'<dLblPos val="(\w+)"', dl), fmt=re.findall(r'<numFmt formatCode="([^"]+)"', dl),
            show={k: re.search(rf'<{k} val="(\d)"', dl).group(1) if re.search(rf'<{k} val="(\d)"', dl) else None
                  for k in ("showVal", "showSerName", "showCatName", "showLegendKey", "showPercent")},
            font=re.search(r'<defRPr sz="(\d+)" b="(\d)"><solidFill><srgbClr val="(\w+)"', dl)))
    want_series = [("'Q2_Summary'!$B$50", "'Q2_Summary'!$B$51:$B$56", "0072B2"),
                   ("'Q2_Summary'!$C$50", "'Q2_Summary'!$C$51:$C$56", "009E73")]
    rec("Chart XML (Figure 2)", "exactly two series", "", "", len(chart_lines), 2, len(chart_lines) == 2)
    for g, (tx, val, colour) in zip(chart_lines, want_series):
        font_ok = g["font"] is not None and g["font"].groups() == ("800", "1", colour)
        ok = (g["tx"] == tx and g["val"] == val and g["cat"] == "'Q2_Summary'!$A$51:$A$56" and g["fill"] == colour
              and g["pos"] == ["outEnd"] and g["fmt"] == ["0.0%"] and font_ok
              and g["show"] == {"showVal": "1", "showSerName": "0", "showCatName": "0", "showLegendKey": "0",
                                "showPercent": "0"})
        rec("Chart XML (Figure 2)", f"series {tx}", "", "",
            f"val={g['val']} cat={g['cat']} fill=#{g['fill']} no outline; labels {g['pos']} {g['fmt']} "
            f"{'8 pt bold #' + colour if font_ok else 'font mismatch'}; value only={g['show']['showVal'] == '1'}",
            f"{val}, #{colour}, outEnd 0.0%, 8 pt bold", ok)
    asserts = {
        "barChart with barDir col": "<barChart>" in xml and '<barDir val="col"/>' in xml,
        "grouping clustered": '<grouping val="clustered"/>' in xml,
        "gapWidth 40": '<gapWidth val="40"/>' in xml,
        "overlap -12": '<overlap val="-12"/>' in xml,
        "no lineChart element": "<lineChart" not in xml,
        "no other chart types in the plot area": not re.search(
            r"<(area|pie|doughnut|scatter|radar|bubble|stock|surface|bar3D|line3D)Chart", xml),
        "no majorGridlines element": "majorGridlines" not in xml,
        "value axis deleted (hidden)": '<delete val="1"/>' in val_ax,
        "value axis min 0, max 0.8": '<min val="0"/>' in val_ax and '<max val="0.8"/>' in val_ax,
        "value axis major unit 0.2, format 0%": '<majorUnit val="0.2"/>' in val_ax and 'formatCode="0%"' in val_ax,
        "category axis shown (delete 0)": '<delete val="0"/>' in cat_ax,
        "category axis tickLblSkip 1, tickMarkSkip 1": '<tickLblSkip val="1"/>' in cat_ax
                                                         and '<tickMarkSkip val="1"/>' in cat_ax,
        "category labels rot 0": '<txPr><bodyPr rot="0"' in cat_ax,
        "category labels Calibri 8 pt #595959": re.search(r'<txPr><bodyPr rot="0"[^>]*/><p><pPr><defRPr sz="800" '
                                                          r'b="0"><solidFill><srgbClr val="595959"', cat_ax) is not None,
        "category axis line #BFBFBF 0.75 pt": '<ln w="9525"><solidFill><srgbClr val="BFBFBF"/>' in cat_ax,
        "axis title 'Booking lead time (days before arrival)'":
            "<t>Booking lead time (days before arrival)</t>" in cat_ax,
        "legend at top (legendPos t)": '<legendPos val="t"/>' in legend,
        "legend Calibri 9 pt #404040": re.search(r'<defRPr sz="900" b="0"><solidFill><srgbClr val="404040"',
                                                 legend) is not None,
        "no chart title (autoTitleDeleted 1)": '<autoTitleDeleted val="1"/>' in xml and "<chart><title>" not in xml,
        "chart area white, no border": "<spPr><solidFill><srgbClr val=\"FFFFFF\"/></solidFill><ln><noFill/>" in xml,
        "no 3-D": "3D" not in xml,
        "Calibri throughout": 'typeface="Calibri"' in xml and "typeface=" not in xml.replace('typeface="Calibri"', ""),
    }
    for k, v in asserts.items():
        rec("Chart XML (Figure 2)", k, "", "", v, True, v)
    drawing = [n for n in z.namelist() if n.startswith("xl/drawings/drawing")]
    dx = re.sub(r"<(/?)[a-zA-Z0-9]+:", r"<\1", z.read(drawing[0]).decode("utf8"))
    ext = re.search(r'<ext cx="(\d+)" cy="(\d+)"', dx)
    w_cm, h_cm = int(ext.group(1)) / 360000, int(ext.group(2)) / 360000
    rec("Chart XML (Figure 2)", "size (cm)", "", "", f"{w_cm:.2f} × {h_cm:.2f}", "15.9 × 9.5",
        abs(w_cm - 15.9) < 0.01 and abs(h_cm - 9.5) < 0.01)
    frm = re.search(r"<from><col>(\d+)</col><colOff>\d+</colOff><row>(\d+)</row>", dx)
    rec("Chart XML (Figure 2)", "anchor", "", "", f"col {frm.group(1)}, row {frm.group(2)} (0-based)", "Figure2!A4",
        frm.groups() == ("0", "3"))
    wbx = z.read("xl/workbook.xml").decode("utf8")
    fig_sheet_idx = re.findall(r'<sheet [^>]*name="([^"]+)"', re.sub(r"<(/?)[a-zA-Z0-9]+:", r"<\1", wbx)).index("Figure2")
    rels = z.read(f"xl/worksheets/_rels/sheet{fig_sheet_idx + 1}.xml.rels").decode("utf8")
    rec("Chart XML (Figure 2)", "chart sits on the Figure2 sheet", "", "", "drawing rel present", True,
        "drawing" in rels)
    chart_assert_list = [(k, v) for k, v in asserts.items()]

    rc_text = open(os.path.join(os.path.dirname(recalc), "recalc_time.txt")).read()
    rc_json = json.loads(rc_text[rc_text.index("{"):rc_text.rindex("}") + 1])
    rec("Structure", "LibreOffice recalculation: formula error values", "", "",
        f"{rc_json.get('total_errors')} errors in {rc_json.get('total_formulas'):,} formulas", 0,
        rc_json.get("status") == "success" and rc_json.get("total_errors") == 0)

    size_mb = os.path.getsize(built) / 1e6
    recalc_time = open(os.path.join(os.path.dirname(recalc), "recalc_time.txt")).read().split("\n")[0]

    # ---- write report
    n_fail = sum(1 for r in results if r[-1] == "FAIL")
    out = [f"# Verification report v3: {os.path.basename(built)}", "",
           f"Generated by `verify.py` on {datetime.now():%d %B %Y %H:%M}.", "",
           f"**Overall: {'ALL PASS' if n_fail == 0 else f'{n_fail} FAIL'}** "
           f"({len(results) - n_fail} of {len(results)} checks passed).", "",
           "- Method A = pandas on the original file; Method B = plain openpyxl row loop on the original file "
           "(explicit if/elif bands, Python set for duplicates); Method C = the workbook's Excel formulas, "
           "recalculated by LibreOffice and read with `data_only=True`.",
           "- Counts must agree exactly; rates to within 1e-12. Expected values (brief Section 7) are compared after "
           "Decimal half-up rounding to 1 decimal place in per cent (all three methods must agree), and must sit within "
           "±0.005 percentage points of the 2-dp golden figure.",
           "", "## Build facts", "",
           f"- Workbook size: {size_mb:.1f} MB",
           f"- LibreOffice recalculation (headless, includes load, full recalculation, save and an error scan): "
           f"{recalc_time} s",
           f"- Verification run time: {time.time() - t_start:.0f} s", ""]
    out += ["## Change log, v1 to v3", "",
            "- **Figure 2 redesigned** as a native clustered column chart (openpyxl BarChart, built with "
            "`build_figure2()` from `figure2_final_reference.py`): two series (all bookings #0072B2, no-deposit "
            "bookings #009E73), 0.0% value labels on all 12 bars, value axis hidden on a fixed 0–80% scale, "
            "legend at the top, 15.9 × 9.5 cm. The v1 line chart, its online TA series and the label nudges are gone.",
            "- **Table D (Q2_Summary rows 49–57) rebuilt**: three-line category labels (band / bookings / "
            "cancelled; Over 365 adds the non-refundable share), the two rate series and a 'Check: equals Table "
            "A?' column. Table A keeps its online TA columns, which the report text cites.",
            "- **Figure2 sheet text**: new caption, note and source lines (number-bearing text is formulas); print "
            "area A1:J27, A4 portrait, fit to one page.",
            "- **Checks**: the two online TA chart checks were replaced by six figure checks (Table D equals "
            "Table A, every label shows bookings and cancellations, the Over 365 label shows the non-refundable "
            "share, every value fits the 80% scale, cancel_flag has no errors).",
            "- **README**: Figure 2 sheet-map line, suppression rule, style guide (legend top, chart type, labels "
            "instead of axis, context in labels), a Figure 2 design section, version and open-and-save lines.",
            "- **Unchanged from v1 (as requested)**: Raw_Data, Clean_Data, Table1_DataQuality, Tables A–C, "
            "Report_Figures and every formula outside Table D.", "",
            "## Chart series references (parsed from xl/charts XML)", "",
            "| Series name ref | Values | Categories | Fill | Data labels |",
            "|---|---|---|---|---|"]
    for g in chart_lines:
        out.append(f"| `{g['tx']}` | `{g['val']}` | `{g['cat']}` | #{g['fill']} | {', '.join(g['pos'])}, "
                   f"{', '.join(g['fmt'])} |")
    out += ["", "## Chart XML assertion list", ""]
    out += [f"- {'PASS' if v else 'FAIL'}: {k}" for k, v in chart_assert_list]
    out += ["", f"Rendered preview: `Figure2_v3_preview.png` (LibreOffice, 220 dpi; checking only, not for the report).",
            ""]
    section = None
    for sec, name, a, b, c, exp, st in results:
        if sec != section:
            out += ["", f"## {sec}", "", "| Check | Method A (pandas) | Method B (openpyxl loop) | Method C (Excel formulas) | Expected | Result |",
                    "|---|---|---|---|---|---|"]
            section = sec
        esc = lambda x: str(x).replace("|", "\\|").replace("\n", "⏎")
        out.append(f"| {esc(name)} | {esc(a)} | {esc(b)} | {esc(c)} | {esc(exp)} | {st} |")
    out += ["", "## Deviations and notes", "",
            "- The source file stores missing values as the text `NULL` (country 488, agent 16,340, company 112,593) "
            "and `NA` (children 4). pandas reads these as NaN. As the brief requires, Raw_Data and Clean_Data hold "
            "them as empty cells (129,425 cells). This is the only difference from the source; every other cell "
            "matches exactly (see Structure).",
            "- No-deposit Online TA total cancelled = 20,668, confirmed by all three methods.",
            "- Rates are shown as the 1-dp half-up value with the exact per-cent value in brackets. Golden values "
            "are given to 2 dp, so a rate passes when the three methods agree at 1 dp and sit within ±0.005 "
            "percentage points of the golden figure. One golden value is a rounding tie: Over 365 no-deposit "
            "rate = 264 / 1,280 = 20.625% exactly, listed as 20.62% in the brief (half-up would give 20.63%); at "
            "1 dp both are 20.6%, the figure the report quotes.",
            "- Table 1's last row (186) adds four counts; the distinct number of bookings involved is also 186 "
            "(no overlap), confirmed by all three methods.",
            "- Table 1 text cells reproduce the report wording with live numbers; the only character difference is "
            "a hyphen-minus in \"-6.38\" where the report uses the typographic minus sign (−).",
            "- Visual QA previews were rendered by LibreOffice (Carlito, metric-compatible with Calibri), whose "
            "chart layout differs slightly from Excel's.",
            "- The recalculation time above includes openpyxl's scan of all formulas for error values; LibreOffice "
            "itself loads, recalculates and saves the workbook in about 75 s.",
            "- Method C uses the xlsx skill's recalc.py (LibreOffice headless, an explicit calculateAll() and save), "
            "which is stricter than a plain --convert-to xlsx because it forces every formula to recalculate.",
            ""]
    with open(report, "w") as fh:
        fh.write("\n".join(out))
    print(f"{len(results) - n_fail} / {len(results)} PASS; report written to {report}")
    for r in results:
        if r[-1] == "FAIL":
            print("FAIL:", r)
    return n_fail


# ================================================================== stage: render
def stage_render(recalc, out_dir):
    work = os.path.join(os.path.dirname(recalc), "render")
    os.makedirs(work, exist_ok=True)
    t = time.time()
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", work, recalc], check=True,
                   capture_output=True, timeout=3600)
    pdf = os.path.join(work, os.path.splitext(os.path.basename(recalc))[0] + ".pdf")
    print(f"PDF rendered in {time.time() - t:.0f}s")
    n_pages = int(re.search(r"Pages:\s+(\d+)", subprocess.run(["pdfinfo", pdf], capture_output=True,
                                                              text=True).stdout).group(1))
    found = {}
    for p in range(1, n_pages + 1):
        txt = subprocess.run(["pdftotext", "-f", str(p), "-l", str(p), pdf, "-"], capture_output=True,
                             text=True).stdout
        if "Figure 2: Cancellation rate by booking lead time" in txt and "Figure2" not in found:
            found["Figure2"] = p
        if "Q2: How does booking lead time affect" in txt:
            found.setdefault("Q2", []).append(p)
    print("pages:", n_pages, found)
    q = [found["Q2"][0], found["Figure2"] - 1]  # Q2_Summary runs until the Figure2 sheet starts
    # LibreOffice draws a cached #N/A as 0 in a line chart, whereas Excel leaves the point out. For the Figure 2
    # preview only, blank the suppressed (#N/A) chart-data cells in a throwaway copy so the image matches Excel.
    from openpyxl import load_workbook
    wb = load_workbook(recalc, read_only=True, data_only=True)
    ws = wb["Q2_Summary"]
    na_cells = [c.coordinate for row in ws.iter_rows(min_row=1, max_row=80, max_col=6) for c in row
                if c.value == "#N/A"]
    sheet_xml = f"xl/worksheets/sheet{wb.sheetnames.index('Q2_Summary') + 1}.xml"
    wb.close()
    print("blanked for the Figure 2 preview only:", na_cells)
    copy = os.path.join(work, "figure2_preview_copy.xlsx")
    with zipfile.ZipFile(recalc) as zin, zipfile.ZipFile(copy, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == sheet_xml:
                xml = data.decode("utf8")
                for ref in na_cells:
                    xml, n = re.subn(rf'<c r="{ref}"([^>]*?) t="e"([^>]*)>.*?</c>', rf'<c r="{ref}"\1\2/>', xml,
                                     flags=re.S)
                    assert n == 1, ref
                data = xml.encode("utf8")
            if item.filename.startswith("xl/charts/chart"):
                # LibreOffice's own re-save turns dispBlanksAs "gap" into "zero"; the delivered file keeps "gap"
                data = data.replace(b'dispBlanksAs val="zero"', b'dispBlanksAs val="gap"')
            zout.writestr(item, data)
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", work, copy], check=True,
                   capture_output=True, timeout=3600)
    fpdf = os.path.splitext(copy)[0] + ".pdf"
    subprocess.run(["pdftoppm", "-png", "-r", "220", "-f", str(found["Figure2"]), "-l", str(found["Figure2"]),
                    "-singlefile", fpdf, os.path.join(out_dir, "Figure2_v3_preview")], check=True)
    subprocess.run(["pdfseparate", "-f", str(q[0]), "-l", str(q[-1]), pdf, os.path.join(work, "q2_%d.pdf")],
                   check=True)
    parts = [os.path.join(work, f"q2_{i}.pdf") for i in range(q[0], q[-1] + 1)]
    if len(parts) == 1:
        shutil.copy(parts[0], os.path.join(out_dir, "Q2_Summary_v3_preview.pdf"))
    else:
        subprocess.run(["pdfunite", *parts, os.path.join(out_dir, "Q2_Summary_v3_preview.pdf")], check=True)
    subprocess.run(["pdftoppm", "-png", "-r", "110", "-f", "1", "-l", "1", "-singlefile",
                    os.path.join(out_dir, "Q2_Summary_v3_preview.pdf"), os.path.join(work, "Q2_Summary_page1")],
                   check=True)


if __name__ == "__main__":
    stage = sys.argv[1]
    if stage == "recalc":
        stage_recalc(*sys.argv[2:5])
    elif stage == "check":
        sys.exit(1 if stage_check(*sys.argv[2:6]) else 0)
    elif stage == "render":
        stage_render(*sys.argv[2:4])
